import copy, json
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from lxml import etree
from xml.sax.saxutils import escape

A='http://schemas.openxmlformats.org/drawingml/2006/main'
NS={'a':A,'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
IN=914400
X0=270000; W0=9108000

def rpr(sz,b,color):
    return (f'<a:rPr lang="ko-KR" sz="{sz}" b="{1 if b else 0}"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="Pretendard GOV"/><a:ea typeface="Pretendard GOV"/><a:cs typeface="Pretendard GOV"/></a:rPr>')

def para(text,sz,b,color,algn='l'):
    # text: str 또는 [(str,bold,color)] 조각 목록. 줄바꿈은 문단 분리
    out=''
    lines=text if isinstance(text,list) and text and isinstance(text[0],list) else None
    if isinstance(text,str):
        for ln in text.split('\n'):
            out+=f'<a:p><a:pPr algn="{algn}"/><a:r>{rpr(sz,b,color)}<a:t>{escape(ln)}</a:t></a:r><a:endParaRPr lang="ko-KR" sz="{sz}"/></a:p>'
        return out
    # 조각 목록 한 문단
    runs=''.join(f'<a:r>{rpr(sz,bb if bb is not None else b,cc or color)}<a:t>{escape(t)}</a:t></a:r>' for t,bb,cc in text)
    return f'<a:p><a:pPr algn="{algn}"/>{runs}<a:endParaRPr lang="ko-KR" sz="{sz}"/></a:p>'

LN=''.join(f'<a:{s} w="6350" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:srgbClr val="BFC4D6"/></a:solidFill><a:prstDash val="solid"/></a:{s}>' for s in ('lnL','lnR','lnT','lnB'))

def tc(text,head,fill=None,algn='l',b=None,span=None,hmerge=False):
    sz=800
    color='1F2A55' if head else '404040'
    bold=head if b is None else b
    fl=fill or ('D9DEEC' if head else 'FFFFFF')
    attr=''
    if span: attr+=f' gridSpan="{span}"'
    if hmerge: attr+=' hMerge="1"'
    return (f'<a:tc{attr}><a:txBody><a:bodyPr wrap="square"/><a:lstStyle/>{para(text,sz,bold,color,algn)}</a:txBody>'
            f'<a:tcPr marL="36000" marR="36000" marT="36000" marB="36000" anchor="ctr">{LN}<a:solidFill><a:srgbClr val="{fl}"/></a:solidFill></a:tcPr></a:tc>')

_id=[5000]
def nid():
    _id[0]+=1; return _id[0]

def table(slide,x,y,w,ratios,rows,head_rows=1,aligns=None,fills=None,rowh=188441,bold_col0=False):
    tot=sum(ratios); ws=[int(w*r/tot) for r in ratios]; ws[-1]=w-sum(ws[:-1])
    grid=''.join(f'<a:gridCol w="{c}"/>' for c in ws)
    trs=''
    for i,row in enumerate(rows):
        head=i<head_rows
        cells=''
        j=0
        for cell in row:
            span=None; f=None; txt=cell
            if isinstance(cell,dict):
                txt=cell['t']; span=cell.get('span'); f=cell.get('fill')
            al='ctr' if head else ((aligns[j] if aligns else 'l'))
            if fills and not head and i in fills: f=f or fills[i]
            b=None
            if bold_col0 and j==0 and not head: b=True
            cells+=tc(txt,head,f,al,b,span)
            if span:
                for _ in range(span-1): cells+=tc('',head,f,al,None,None,True)
                j+=span
            else: j+=1
        trs+=f'<a:tr h="{rowh}">{cells}</a:tr>'
    i=nid()
    xml=(f'<p:graphicFrame xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:nvGraphicFramePr><p:cNvPr id="{i}" name="Table {i}"/><p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
         f'<p:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{rowh*len(rows)}"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
         f'<a:tbl><a:tblPr firstRow="0" bandRow="0"/><a:tblGrid>{grid}</a:tblGrid>{trs}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')
    slide.shapes._spTree.append(etree.fromstring(xml))

def bar(slide,x,y,w,text,h=223200):
    i=nid()
    xml=(f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="{i}" name="Bar {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
         f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="2E4692"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>'
         f'<p:txBody><a:bodyPr rtlCol="0" anchor="ctr" lIns="90000" rIns="36000" tIns="36000" bIns="36000" wrap="square"/><a:lstStyle/>{para(text,1050,True,"FFFFFF")}</p:txBody></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))

def sub(slide,x,y,w,text,h=200000):
    # 소제목: 연한 바탕 + 진한 글씨 (레퍼 '예외·운영 규칙' 형식)
    i=nid()
    xml=(f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="{i}" name="Sub {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
         f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="D9DEEC"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr>'
         f'<p:txBody><a:bodyPr rtlCol="0" anchor="ctr" lIns="72000" rIns="36000" tIns="18000" bIns="18000" wrap="square"/><a:lstStyle/>{para(text,900,True,"1F2A55")}</p:txBody></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))

def text(slide,x,y,w,h,content,sz=1000,color='404040',b=False,bullets=False):
    i=nid()
    ps=''
    for ln in content.split('\n') if isinstance(content,str) else content:
        if isinstance(ln,list):
            runs=''.join(f'<a:r>{rpr(sz,bb,cc or color)}<a:t>{escape(t)}</a:t></a:r>' for t,bb,cc in ln)
        else:
            runs=f'<a:r>{rpr(sz,b,color)}<a:t>{escape(ln)}</a:t></a:r>'
        bu='<a:buFont typeface="Arial"/><a:buChar char="•"/>' if bullets else '<a:buNone/>'
        ind=' marL="114300" indent="-114300"' if bullets else ''
        ps+=f'<a:p><a:pPr algn="l"{ind}><a:spcAft><a:spcPts val="200"/></a:spcAft>{bu}</a:pPr>{runs}</a:p>'
    xml=(f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="{i}" name="Text {i}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
         f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
         f'<p:txBody><a:bodyPr wrap="square" lIns="36000" rIns="36000" tIns="18000" bIns="18000"/><a:lstStyle/>{ps}</p:txBody></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))

def box(slide,x,y,w,h,title,body,fill='E9EAF1'):
    i=nid()
    ps=para(title,1000,True,'1F2A55','ctr')+para(body,800,False,'404040','ctr')
    xml=(f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="{i}" name="Box {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
         f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 12000"/></a:avLst></a:prstGeom><a:solidFill><a:srgbClr val="{fill}"/></a:solidFill><a:ln w="9525"><a:solidFill><a:srgbClr val="2E4692"/></a:solidFill></a:ln></p:spPr>'
         f'<p:txBody><a:bodyPr rtlCol="0" anchor="ctr" lIns="36000" rIns="36000" tIns="36000" bIns="36000" wrap="square"/><a:lstStyle/>{ps}</p:txBody></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))

def arrow(slide,x,y,w=111600,h=129600):
    i=nid()
    xml=(f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="{i}" name="Arrow {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
         f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="rightArrow"><a:avLst/></a:prstGeom><a:solidFill><a:srgbClr val="7F7F7F"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
    slide.shapes._spTree.append(etree.fromstring(xml))

def walk(shapes):
    for sh in shapes:
        if sh.shape_type==6: yield from walk(sh.shapes)
        else: yield sh

def set_run_text(sh,new):
    # 첫 run 서식을 살리고 글자만 바꾼다
    tf=sh.text_frame
    p=tf.paragraphs[0]
    if p.runs:
        p.runs[0].text=new
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
    for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)

def frame(slide,title,kind,notes):
    for sh in walk(slide.shapes):
        if not sh.has_text_frame: continue
        t=sh.text_frame.text
        if sh.name=='직사각형 3': set_run_text(sh,title)
        elif t.startswith('□') or t.startswith('■'): set_run_text(sh,kind)
        elif t=='안승민': set_run_text(sh,'')
    fs=[sh for sh in slide.shapes if sh.name=='Group 1019'][0]
    for r in range(10):
        cell=fs.table.cell(r,1)
        p=cell.text_frame.paragraphs[0]._p
        for rr in p.findall('{%s}r'%A): p.remove(rr)
        if r < len(notes):
            p.append(etree.fromstring(f'<a:r xmlns:a="{A}"><a:rPr lang="ko-KR" sz="700" b="0"><a:solidFill><a:srgbClr val="262626"/></a:solidFill><a:latin typeface="맑은 고딕"/><a:ea typeface="맑은 고딕"/><a:cs typeface="맑은 고딕"/></a:rPr><a:t>{escape(notes[r])}</a:t></a:r>'))

def _sp(x,y,w,h,geom,fill,ln=None,body=''):
    i=nid()
    l=f'<a:ln w="{ln[1]}"><a:solidFill><a:srgbClr val="{ln[0]}"/></a:solidFill></a:ln>' if ln else '<a:ln><a:noFill/></a:ln>'
    f=f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else '<a:noFill/>'
    tb=body or '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="ko-KR"/></a:p></p:txBody>'
    return (f'<p:sp xmlns:p="{NS["p"]}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="{i}" name="Shape {i}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm><a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>{f}{l}</p:spPr>{tb}</p:sp>')

def _lines(t,w,sz):
    # 한 줄에 들어가는 글자 수 대략 계산 (한글 폭 = 글자 크기)
    per=max(10,int((w/12700)/(sz/100)*1.35))
    n=0
    for ln in t.split('\n'): n+=max(1,-(-len(ln)//per))
    return n

def dotlist(slide,x,y,w,items,sz=1000,color='404040',gap=40000):
    """원본의 파란 점 + 글 목록. 끝난 y를 돌려준다"""
    lh=int(sz*12700/100*1.35)
    for t in items:
        n=_lines(t,w-130000,sz); h=lh*n+36000
        slide.shapes._spTree.append(etree.fromstring(_sp(x+20000,y+int(lh*0.5)-12000,57600,57600,'ellipse','2E4692')))
        body=(f'<p:txBody><a:bodyPr wrap="square" lIns="18000" rIns="18000" tIns="18000" bIns="18000"/><a:lstStyle/>'
              f'<a:p><a:pPr algn="l"><a:lnSpc><a:spcPct val="110000"/></a:lnSpc></a:pPr><a:r>{rpr(sz,False,color)}<a:t>{escape(t)}</a:t></a:r></a:p></p:txBody>')
        el=etree.fromstring(_sp(x+130000,y,w-130000,h,'rect',None,None,body)); el.find('.//{%s}cNvSpPr'%NS['p']).set('txBox','1')
        slide.shapes._spTree.append(el)
        y+=h+gap
    return y

def callout(slide,x,y,w,label,items,sz=1000):
    """원본의 노란 안내 상자 (왼쪽 굵은 제목 + 오른쪽 글)"""
    lw=1080000
    lh=int(sz*12700/100*1.35)
    n=sum(_lines(t,w-lw-200000,sz) for t in items)
    h=lh*n+30000*len(items)+80000
    slide.shapes._spTree.append(etree.fromstring(_sp(x,y,w,h,'rect','FFF8E1',('E0B000',9525))))
    lb=(f'<p:txBody><a:bodyPr wrap="square" anchor="t" lIns="18000" rIns="18000" tIns="18000" bIns="18000"/><a:lstStyle/>'
        +''.join(f'<a:p><a:pPr algn="l"/><a:r>{rpr(sz,True,"7F6000")}<a:t>{escape(t)}</a:t></a:r></a:p>' for t in label.split('\n'))+'</p:txBody>')
    el=etree.fromstring(_sp(x+72000,y+45000,lw-72000,h-90000,'rect',None,None,lb)); el.find('.//{%s}cNvSpPr'%NS['p']).set('txBox','1')
    slide.shapes._spTree.append(el)
    ps=''.join(f'<a:p><a:pPr algn="l"><a:lnSpc><a:spcPct val="110000"/></a:lnSpc><a:spcAft><a:spcPts val="300"/></a:spcAft></a:pPr><a:r>{rpr(sz,False,"404040")}<a:t>{escape(t)}</a:t></a:r></a:p>' for t in items)
    tb=f'<p:txBody><a:bodyPr wrap="square" anchor="t" lIns="18000" rIns="18000" tIns="18000" bIns="18000"/><a:lstStyle/>{ps}</p:txBody>'
    el=etree.fromstring(_sp(x+lw,y+45000,w-lw-72000,h-90000,'rect',None,None,tb)); el.find('.//{%s}cNvSpPr'%NS['p']).set('txBox','1')
    slide.shapes._spTree.append(el)
    return y+h
