# 열린 PPT에 큰 표를 넣을 때: python-pptx로 임시 pptx에 표를 만들고(WAS-02 서식, 세로 병합 포함)
# paste_table_example.ps1처럼 COM으로 임시 파일을 열어 Copy → 열린 슬라이드에 Paste 한다.
# 칸마다 COM으로 칠하면 290칸에 5분을 넘겨 끊겼다(2026-10-20). 이 방식은 몇 초.
from pptx import Presentation
from lxml import etree
from xml.sax.saxutils import escape
A='http://schemas.openxmlformats.org/drawingml/2006/main'; P='http://schemas.openxmlformats.org/presentationml/2006/main'
LN=''.join(f'<a:{s} w="6350" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:srgbClr val="BFC4D6"/></a:solidFill><a:prstDash val="solid"/></a:{s}>' for s in ('lnL','lnR','lnT','lnB'))

def tc(t, fill, col, bold, sz=900, algn='ctr', extra=''):
    """extra: ' rowSpan="2"'(위 칸) / ' vMerge="1"'(아래 이어진 칸) / ' gridSpan="2"' / ' hMerge="1"'"""
    return (f'<a:tc{extra}><a:txBody><a:bodyPr/><a:lstStyle/><a:p><a:pPr algn="{algn}"/><a:r><a:rPr lang="ko-KR" sz="{sz}" b="{1 if bold else 0}"><a:solidFill><a:srgbClr val="{col}"/></a:solidFill><a:latin typeface="Pretendard GOV"/><a:ea typeface="Pretendard GOV"/><a:cs typeface="Pretendard GOV"/></a:rPr><a:t>{escape(t)}</a:t></a:r><a:endParaRPr lang="ko-KR" sz="{sz}"/></a:p></a:txBody>'
            f'<a:tcPr marL="36000" marR="36000" marT="36000" marB="36000" anchor="ctr">{LN}<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill></a:tcPr></a:tc>')

def add_table(slide, name, col_widths_emu, rows_xml, x=266700, y=1155700, row_h=165100):
    """rows_xml: 행마다 tc() 문자열을 이어 붙인 목록. 머리글 D9DEEC/1F2A55 굵게, 첫 열 F7F8FB, 본문 FFFFFF/404040"""
    grid=''.join(f'<a:gridCol w="{c}"/>' for c in col_widths_emu)
    trs=''.join(f'<a:tr h="{row_h}">{r}</a:tr>' for r in rows_xml)
    xml=(f'<p:graphicFrame xmlns:p="{P}" xmlns:a="{A}"><p:nvGraphicFramePr><p:cNvPr id="900" name="{name}"/><p:cNvGraphicFramePr><a:graphicFrameLocks noGrp="1"/></p:cNvGraphicFramePr><p:nvPr/></p:nvGraphicFramePr>'
         f'<p:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{sum(col_widths_emu)}" cy="{row_h*len(rows_xml)}"/></p:xfrm><a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/table">'
         f'<a:tbl><a:tblPr firstRow="0" bandRow="0"/><a:tblGrid>{grid}</a:tblGrid>{trs}</a:tbl></a:graphicData></a:graphic></p:graphicFrame>')
    slide.shapes._spTree.append(etree.fromstring(xml))

# 사용: prs=Presentation(<원본 복사본>); s=prs.slides.add_slide(prs.slides[n].slide_layout); add_table(s,'BandTable',[...],[...]); prs.save(<임시 경로>)
