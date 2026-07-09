# -*- coding: utf-8 -*-
"""SVC 카테고리(서비스 구조도/사업 흐름도/운영 프로세스도) 에셋 생성 — SVC-001~SVC-008.

- 네이티브 도형/커넥터/화살표만 사용 (SmartArt·이미지 금지).
- 대표 도형에 c.name_asset("SVC-00N"), 좌상단 c.id_caption.
- 색은 c.role/c.C 만 사용 (네이비/블루/틸/그레이 코퍼레이트).
"""
import sys
import os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lib'))
import common as c
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

R = c.role
CC = c.C


# ---------- 공용 헬퍼 ----------
def title_block(slide, text, sub=None):
    c.add_text(slide, 0.55, 0.52, 12.2, 0.5, text, size=18, bold=True,
               color=R("header_fill"), align=PP_ALIGN.LEFT, font=c.FONT_H)
    if sub:
        c.add_text(slide, 0.55, 1.02, 12.2, 0.32, sub, size=11,
                   color=R("muted_text"), align=PP_ALIGN.LEFT)


def node(slide, x, y, w, h, text, fill, tcolor=None, line=None, line_w=1.0,
         size=11, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    box = c.add_box(slide, x, y, w, h, fill=fill, line=line or R("border"),
                    line_w=line_w, shape=shape)
    c.set_shape_text(box, text, size=size, bold=bold,
                     color=tcolor or R("header_text"))
    return box


def arrow(slide, x, y, w, h, direction="right", color=None, text=None, size=10):
    shp = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW,
           "left": MSO_SHAPE.LEFT_ARROW, "up": MSO_SHAPE.UP_ARROW,
           "lr": MSO_SHAPE.LEFT_RIGHT_ARROW}[direction]
    a = c.add_box(slide, x, y, w, h, fill=color or R("accent_point"),
                  line=None, shape=shp)
    if text:
        c.set_shape_text(a, text, size=size, bold=True, color=R("header_text"))
    return a


def band_label(slide, x, y, w, h, text, fill, size=12):
    lb = c.add_box(slide, x, y, w, h, fill=fill, line=None,
                   shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    c.set_shape_text(lb, text, size=size, bold=True, color=R("header_text"))
    return lb


# =========================================================
# SVC-001  시스템 구성도 (수집→처리→저장→서비스 4계층, 세로)
# =========================================================
def system_layers(slide, asset_id):
    title_block(slide, "시스템 구성도 (4계층 아키텍처)",
                "데이터 수집 → 정제·가공 → DB 적재 → 서비스 제공 (계층 간 하향 흐름)")
    x0, xw = 0.9, 11.5
    lab_w = 2.3
    item_x = x0 + lab_w + 0.35
    item_w = xw - lab_w - 0.35
    y = 1.55
    lh = 1.05
    gap = 0.38
    layers = [
        ("수집 계층\n데이터 수집", CC["navy_800"],
         [["공공데이터 API", "내부 시스템", "외부 데이터원"], R("accent_primary")]),
        ("처리 계층\n정제·가공", CC["navy_600"],
         [["데이터 정제", "표준화·변환", "품질 검증"], R("accent_point")]),
        ("저장 계층\nDB 적재", R("accent_secondary"),
         [["운영 DB", "데이터 웨어하우스", "메타데이터"], R("accent_secondary")]),
        ("서비스 계층\nAPI 제공", R("sub_header"),
         [["Open API", "관리 대시보드", "통계 리포트"], R("sub_header")]),
    ]
    cont = c.add_box(slide, x0 - 0.05, y - 0.05, xw + 0.1,
                     (lh + gap) * len(layers) - gap + 0.1, fill=None, line=None)
    c.name_asset(cont, asset_id)
    for i, (lab, lfill, (items, icol)) in enumerate(layers):
        ly = y + i * (lh + gap)
        band_label(slide, x0, ly, lab_w, lh, lab, lfill)
        n = len(items)
        g = 0.25
        bw = (item_w - g * (n - 1)) / n
        for j, it in enumerate(items):
            node(slide, item_x + j * (bw + g), ly, bw, lh, it,
                 fill=CC["gray_050"], tcolor=R("body_text"),
                 line=icol, line_w=1.5, size=11)
        if i < len(layers) - 1:
            arrow(slide, x0 + lab_w / 2 - 0.22, ly + lh + 0.02, 0.44, gap - 0.04,
                  "down", color=R("accent_point"))
            arrow(slide, item_x + item_w / 2 - 0.22, ly + lh + 0.02, 0.44, gap - 0.04,
                  "down", color=R("accent_point"))


# =========================================================
# SVC-002  3티어 아키텍처 (수평 밴드)
# =========================================================
def three_tier(slide, asset_id):
    title_block(slide, "3-Tier 아키텍처 구조도",
                "사용자 계층 · 애플리케이션 계층 · 데이터 계층 수평 밴드")
    x0, xw = 0.9, 11.5
    y = 1.6
    bh = 1.55
    gap = 0.5
    tiers = [
        ("사용자 계층", CC["navy_800"],
         ["웹 포털", "모바일 앱", "관리자 콘솔"], R("accent_primary")),
        ("애플리케이션 계층", R("accent_primary"),
         ["인증·권한", "업무 로직", "API 게이트웨이", "배치 처리"], R("accent_point")),
        ("데이터 계층", R("accent_secondary"),
         ["운영 DB", "캐시", "파일 스토리지"], R("accent_secondary")),
    ]
    cont = c.add_box(slide, x0 - 0.05, y - 0.05, xw + 0.1,
                     (bh + gap) * len(tiers) - gap + 0.1, fill=None, line=None)
    c.name_asset(cont, asset_id)
    for i, (name, lfill, items, icol) in enumerate(tiers):
        by = y + i * (bh + gap)
        c.add_box(slide, x0, by, xw, bh, fill=CC["gray_050"],
                  line=icol, line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        band_label(slide, x0 + 0.15, by + 0.2, 2.6, bh - 0.4, name, lfill, size=13)
        n = len(items)
        area_x = x0 + 3.0
        area_w = xw - 3.2
        g = 0.22
        bw = (area_w - g * (n - 1)) / n
        for j, it in enumerate(items):
            node(slide, area_x + j * (bw + g), by + 0.32, bw, bh - 0.64, it,
                 fill=R("row_base"), tcolor=R("body_text"), line=icol, line_w=1.0,
                 size=11)
        if i < len(tiers) - 1:
            arrow(slide, x0 + xw / 2 - 0.3, by + bh + 0.02, 0.6, gap - 0.04,
                  "lr", color=R("muted_text"), text="")


# =========================================================
# SVC-003  플랫폼 구조도 (중앙 플랫폼 + 좌우 입력/출력 모듈)
# =========================================================
def platform_hub(slide, asset_id):
    title_block(slide, "플랫폼 구조도 (입력 · 플랫폼 · 출력)",
                "좌측 입력 모듈 → 중앙 통합 플랫폼 → 우측 출력 모듈")
    y0 = 1.7
    mh = 1.15
    mgap = 0.35
    col_w = 3.1
    lx = 0.9
    rx = 13.333 - 0.9 - col_w
    # 중앙 플랫폼
    pw = 3.5
    px = (13.333 - pw) / 2
    py = 1.75
    ph = 4.7
    plat = c.add_box(slide, px, py, pw, ph, fill=CC["navy_800"], line=R("accent_primary"),
                     line_w=2.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    c.name_asset(plat, asset_id)
    c.add_text(slide, px, py + 0.25, pw, 0.5, "통합 데이터 플랫폼", size=15, bold=True,
               color=R("header_text"), align=PP_ALIGN.CENTER, font=c.FONT_H)
    core = ["수집 엔진", "처리 엔진", "저장소", "API 서버"]
    for k, cm in enumerate(core):
        node(slide, px + 0.3, py + 1.0 + k * 0.85, pw - 0.6, 0.68, cm,
             fill=R("accent_primary"), tcolor=R("header_text"), line=None, size=11)
    inputs = ["공공데이터", "내부 시스템", "외부 API"]
    outputs = ["Open API", "대시보드", "성과 리포트"]
    for k, im in enumerate(inputs):
        iy = y0 + k * (mh + mgap)
        node(slide, lx, iy, col_w, mh, im, fill=CC["gray_050"], tcolor=R("body_text"),
             line=R("accent_point"), line_w=1.5)
        arrow(slide, lx + col_w + 0.05, iy + mh / 2 - 0.2, px - (lx + col_w) - 0.1, 0.4,
              "right", color=R("accent_point"))
    for k, om in enumerate(outputs):
        oy = y0 + k * (mh + mgap)
        node(slide, rx, oy, col_w, mh, om, fill=CC["gray_050"], tcolor=R("body_text"),
             line=R("accent_secondary"), line_w=1.5)
        arrow(slide, px + pw + 0.05, oy + mh / 2 - 0.2, rx - (px + pw) - 0.1, 0.4,
              "right", color=R("accent_secondary"))
    c.add_text(slide, lx, y0 - 0.42, col_w, 0.32, "입력 모듈", size=11, bold=True,
               color=R("muted_text"), align=PP_ALIGN.CENTER)
    c.add_text(slide, rx, y0 - 0.42, col_w, 0.32, "출력 모듈", size=11, bold=True,
               color=R("muted_text"), align=PP_ALIGN.CENTER)


# =========================================================
# SVC-004  데이터 흐름도 (소스 다중 → 통합 → 활용처 다중)
# =========================================================
def data_flow(slide, asset_id):
    title_block(slide, "데이터 흐름도 (수집 → 통합 → 활용)",
                "다중 데이터 소스 → 통합 정제 → 다중 활용처 (좌 → 우)")
    sources = ["공공데이터 포털", "내부 업무 시스템", "외부 제휴 API", "센서·로그"]
    uses = ["정책 의사결정", "대민 서비스", "통계 분석", "연계 기관 제공"]
    sw = 2.7
    sx = 0.9
    sh = 0.95
    sgap = 0.32
    sy0 = 1.65
    cont = c.add_box(slide, sx - 0.05, sy0 - 0.05, 13.333 - 1.8 + 0.1,
                     (sh + sgap) * 4 - sgap + 0.1, fill=None, line=None)
    c.name_asset(cont, asset_id)
    # 중앙 통합
    hub_w, hub_h = 2.5, 3.4
    hub_x = (13.333 - hub_w) / 2
    hub_y = (sy0) + ((sh + sgap) * 4 - sgap - hub_h) / 2
    hub = c.add_box(slide, hub_x, hub_y, hub_w, hub_h, fill=CC["navy_800"],
                    line=R("accent_primary"), line_w=2.0, shape=MSO_SHAPE.HEXAGON)
    c.set_shape_text(hub, "데이터\n통합·정제\n(ETL)", size=13, bold=True,
                     color=R("header_text"))
    ux = 13.333 - 0.9 - sw
    for k, s in enumerate(sources):
        yy = sy0 + k * (sh + sgap)
        node(slide, sx, yy, sw, sh, s, fill=CC["gray_050"], tcolor=R("body_text"),
             line=R("accent_point"), line_w=1.5)
        c.connector(slide, sx + sw, yy + sh / 2, hub_x, hub_y + hub_h / 2,
                    color=R("accent_point"), w=1.75)
    for k, u in enumerate(uses):
        yy = sy0 + k * (sh + sgap)
        node(slide, ux, yy, sw, sh, u, fill=CC["gray_050"], tcolor=R("body_text"),
             line=R("accent_secondary"), line_w=1.5)
        c.connector(slide, hub_x + hub_w, hub_y + hub_h / 2, ux, yy + sh / 2,
                    color=R("accent_secondary"), w=1.75)
    # 방향 화살표(중앙 상하)
    arrow(slide, hub_x + hub_w / 2 - 0.9, hub_y - 0.55, 1.8, 0.4, "right",
          color=R("accent_primary"), text="통합")
    arrow(slide, hub_x + hub_w / 2 - 0.9, hub_y + hub_h + 0.15, 1.8, 0.4, "right",
          color=R("accent_primary"), text="활용")


# =========================================================
# SVC-005  서비스 운영 프로세스 (스윔레인 2행)
# =========================================================
def operation_swimlane(slide, asset_id):
    title_block(slide, "서비스 운영 프로세스 (스윔레인)",
                "신청 → 검토 → 승인 → 제공 (신청 기관 · 수행 기관 2개 레인)")
    x0 = 0.9
    lab_w = 1.9
    area_x = x0 + lab_w
    area_w = 13.333 - 0.9 - area_x
    y0 = 1.7
    lane_h = 2.2
    lane_gap = 0.3
    lanes = [("신청 기관", R("accent_primary")), ("수행 기관", R("accent_secondary"))]
    cont = c.add_box(slide, x0 - 0.05, y0 - 0.05, (area_x + area_w) - x0 + 0.1,
                     lane_h * 2 + lane_gap + 0.1, fill=None, line=None)
    c.name_asset(cont, asset_id)
    for i, (name, col) in enumerate(lanes):
        ly = y0 + i * (lane_h + lane_gap)
        c.add_box(slide, area_x, ly, area_w, lane_h, fill=CC["gray_050"],
                  line=R("border"), line_w=1.0)
        band_label(slide, x0, ly, lab_w, lane_h, name, col, size=12)
    # 4 컬럼 스테이지
    ncol = 4
    cgap = 0.55
    sw = (area_w - cgap * (ncol - 1) - 0.6) / ncol
    sh = 1.1
    col_x = [area_x + 0.3 + k * (sw + cgap) for k in range(ncol)]
    top_y = y0 + (lane_h - sh) / 2
    bot_y = y0 + lane_h + lane_gap + (lane_h - sh) / 2
    # 배치: 신청(상,0) → 검토(하,1) → 승인(하,2) → 제공(상,3)
    stages = [
        ("① 신청서 제출", 0, col_x[0], top_y, R("accent_primary")),
        ("② 서류 검토", 1, col_x[1], bot_y, R("accent_secondary")),
        ("③ 승인 심의", 1, col_x[2], bot_y, R("accent_secondary")),
        ("④ 서비스 제공", 0, col_x[3], top_y, R("accent_primary")),
    ]
    centers = []
    for txt, lane, sx, sy, col in stages:
        node(slide, sx, sy, sw, sh, txt, fill=R("row_base"), tcolor=R("body_text"),
             line=col, line_w=1.75, size=12)
        centers.append((sx, sy, col))
    # 화살표 연결 (스테이지 간, 레인 교차 시 대각 커넥터)
    for k in range(ncol - 1):
        sx, sy, col = centers[k]
        nx, ny, ncol_c = centers[k + 1]
        c.connector(slide, sx + sw, sy + sh / 2, nx, ny + sh / 2,
                    color=R("accent_point"), w=2.0)
    # 좌→우 진행 표시 화살표(하단)
    arrow(slide, area_x, y0 + lane_h * 2 + lane_gap + 0.12, area_w, 0.34, "right",
          color=R("muted_text"), text="처리 흐름 (신청 → 제공)")


# =========================================================
# SVC-006  사업 추진 개념도 (로직모델: 입력→활동→산출→성과)
# =========================================================
def logic_model(slide, asset_id):
    title_block(slide, "사업 추진 개념도 (로직 모델)",
                "투입(Input) → 활동(Activity) → 산출(Output) → 성과(Outcome)")
    cols = [
        ("투입", "Input", CC["navy_800"], R("accent_primary"),
         ["예산·인력", "인프라", "데이터 자산"]),
        ("활동", "Activity", CC["navy_600"], R("accent_point"),
         ["시스템 구축", "데이터 정제", "교육·홍보"]),
        ("산출", "Output", R("accent_secondary"), R("accent_secondary"),
         ["구축 시스템", "정제 데이터셋", "운영 매뉴얼"]),
        ("성과", "Outcome", R("sub_header"), R("sub_header"),
         ["업무 효율화", "대민 만족도", "성과 관리 체계"]),
    ]
    n = len(cols)
    x0 = 0.9
    total_w = 13.333 - 1.8
    agap = 0.5
    cw = (total_w - agap * (n - 1)) / n
    y0 = 1.75
    hh = 0.95
    ch = 4.3
    cont = c.add_box(slide, x0 - 0.05, y0 - 0.05, total_w + 0.1, ch + 0.1,
                     fill=None, line=None)
    c.name_asset(cont, asset_id)
    for i, (ko, en, hfill, icol, items) in enumerate(cols):
        cx = x0 + i * (cw + agap)
        c.add_box(slide, cx, y0, cw, ch, fill=CC["gray_050"], line=icol, line_w=1.5,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        hd = c.add_box(slide, cx, y0, cw, hh, fill=hfill, line=None,
                       shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        tf = hd.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = ko; c.set_kfont(r, c.FONT_H, 16, True, R("header_text"))
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = en; c.set_kfont(r2, c.FONT_B, 9, False, R("header_text"))
        for j, it in enumerate(items):
            iy = y0 + hh + 0.28 + j * 0.95
            bx = c.add_box(slide, cx + 0.22, iy, cw - 0.44, 0.72, fill=R("row_base"),
                           line=icol, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            c.set_shape_text(bx, it, size=11, bold=False, color=R("body_text"))
        if i < n - 1:
            arrow(slide, cx + cw + 0.03, y0 + ch / 2 - 0.28, agap - 0.06, 0.56,
                  "right", color=R("accent_point"))


# =========================================================
# SVC-007  정보 연계도 (허브-스포크)
# =========================================================
def hub_spoke(slide, asset_id):
    title_block(slide, "정보 연계도 (기관 간 데이터 연계)",
                "중앙 연계 허브 ↔ 참여 기관 (허브-스포크, 양방향 연계)")
    import math
    cx, cy = 13.333 / 2, 4.35
    hub_w, hub_h = 2.9, 1.5
    hub = c.add_box(slide, cx - hub_w / 2, cy - hub_h / 2, hub_w, hub_h,
                    fill=CC["navy_800"], line=R("accent_primary"), line_w=2.0,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    c.set_shape_text(hub, "데이터 연계 허브", size=14, bold=True, color=R("header_text"))
    c.name_asset(hub, asset_id)
    spokes = ["중앙부처", "광역 지자체", "기초 지자체", "공공기관", "유관 기관", "민간 협력사"]
    rx, ry = 4.7, 2.55
    nsp = len(spokes)
    sw, sh = 2.3, 0.95
    cols = [R("accent_primary"), R("accent_point"), R("accent_secondary"),
            R("sub_header"), R("accent_primary"), R("accent_secondary")]
    for k, sp in enumerate(spokes):
        ang = math.pi * (0.5 + 2 * k / nsp)  # 위에서 시작, 시계 반대
        nx = cx + rx * math.cos(ang)
        ny = cy - ry * math.sin(ang)
        # 커넥터(허브 중심 ↔ 노드 중심) — 양방향 표시
        c.connector(slide, cx, cy, nx, ny, color=R("muted_text"), w=1.75)
        node(slide, nx - sw / 2, ny - sh / 2, sw, sh, sp, fill=CC["gray_050"],
             tcolor=R("body_text"), line=cols[k], line_w=1.5)


# =========================================================
# SVC-008  통합 관제/모니터링 구조 (수집-분석-대시보드-알림)
# =========================================================
def monitoring(slide, asset_id):
    title_block(slide, "통합 관제·모니터링 구조도",
                "데이터 수집 → 실시간 분석 → 대시보드 → 이상 알림")
    x0 = 0.9
    total_w = 13.333 - 1.8
    # 상단 관제 센터 밴드
    center = c.add_box(slide, x0, 1.6, total_w, 0.9, fill=CC["navy_800"], line=None,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    c.set_shape_text(center, "통합 관제 센터 (24시간 모니터링)", size=14, bold=True,
                     color=R("header_text"))
    c.name_asset(center, asset_id)
    # 파이프라인 4단계
    stages = [
        ("수집", "로그·지표\n메트릭 수집", R("accent_primary")),
        ("분석", "실시간 분석\n임계치 판정", R("accent_point")),
        ("대시보드", "현황 시각화\n지표 모니터", R("accent_secondary")),
        ("알림", "이상 탐지\n경보·통지", R("sub_header")),
    ]
    n = len(stages)
    agap = 0.7
    cw = (total_w - agap * (n - 1)) / n
    y = 3.1
    sh = 2.0
    for i, (ko, desc, col) in enumerate(stages):
        sx = x0 + i * (cw + agap)
        c.add_box(slide, sx, y, cw, sh, fill=CC["gray_050"], line=col, line_w=1.75,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        hd = c.add_box(slide, sx, y, cw, 0.62, fill=col, line=None,
                       shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        c.set_shape_text(hd, ko, size=14, bold=True, color=R("header_text"))
        c.add_text(slide, sx + 0.12, y + 0.72, cw - 0.24, sh - 0.8, desc, size=11,
                   color=R("body_text"), align=PP_ALIGN.CENTER)
        # 관제센터 → 각 단계 하향 화살표
        arrow(slide, sx + cw / 2 - 0.16, 2.52, 0.32, 0.5, "down",
              color=R("accent_point"))
        if i < n - 1:
            arrow(slide, sx + cw + 0.05, y + sh / 2 - 0.26, agap - 0.1, 0.52,
                  "right", color=R("accent_point"))
    # 하단 데이터 소스 밴드
    src = c.add_box(slide, x0, y + sh + 0.45, total_w, 0.8, fill=R("accent_primary"),
                    line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    c.set_shape_text(src, "데이터 소스: 서버 · 네트워크 · 애플리케이션 · 사용자 이벤트",
                     size=12, bold=True, color=R("header_text"))
    arrow(slide, x0 + total_w / 2 - 0.16, y + sh + 0.02, 0.32, 0.42, "up",
          color=R("accent_point"))


# =========================================================
# 빌드
# =========================================================
entries = []


def E(*a, **k):
    entries.append(c.entry(*a, **k))


# ---- FILE 1: 아키텍처 (SVC-001~004) ----
p1 = c.new_deck()

s = c.blank_slide(p1); c.id_caption(s, "SVC-001"); system_layers(s, "SVC-001")
s = c.blank_slide(p1); c.id_caption(s, "SVC-002"); three_tier(s, "SVC-002")
s = c.blank_slide(p1); c.id_caption(s, "SVC-003"); platform_hub(s, "SVC-003")
s = c.blank_slide(p1); c.id_caption(s, "SVC-004"); data_flow(s, "SVC-004")

f1 = c.save_deck(p1, "decks/07_service/SVC_architecture_v1.pptx")

# ---- FILE 2: 운영/흐름 (SVC-005~008) ----
p2 = c.new_deck()

s = c.blank_slide(p2); c.id_caption(s, "SVC-005"); operation_swimlane(s, "SVC-005")
s = c.blank_slide(p2); c.id_caption(s, "SVC-006"); logic_model(s, "SVC-006")
s = c.blank_slide(p2); c.id_caption(s, "SVC-007"); hub_spoke(s, "SVC-007")
s = c.blank_slide(p2); c.id_caption(s, "SVC-008"); monitoring(s, "SVC-008")

f2 = c.save_deck(p2, "decks/07_service/SVC_operation-flow_v1.pptx")

# =========================================================
# 매니페스트 엔트리
# =========================================================
F1 = "decks/07_service/SVC_architecture_v1.pptx"
F2 = "decks/07_service/SVC_operation-flow_v1.pptx"

E("SVC-001", "service", "시스템 구성도 (4계층 아키텍처)", F1, 1,
  ["구조도", "시스템구성", "아키텍처", "4계층", "데이터흐름"],
  {"layout": "vertical-layers", "layers": 4, "flow": "top-down"},
  {"layers": [
      {"name": "수집", "items": ["공공데이터 API", "내부 시스템", "외부 데이터원"]},
      {"name": "처리", "items": ["데이터 정제", "표준화·변환", "품질 검증"]},
      {"name": "저장", "items": ["운영 DB", "데이터 웨어하우스", "메타데이터"]},
      {"name": "서비스", "items": ["Open API", "관리 대시보드", "통계 리포트"]}],
   "flow": "top-down"},
  ["layer_labels", "layer_items", "layer_color", "arrows"],
  recommended_use=["시스템 구성", "서비스 구조", "사업 흐름"])

E("SVC-002", "service", "3-Tier 아키텍처 구조도", F1, 2,
  ["구조도", "아키텍처", "3티어", "시스템구성", "수평밴드"],
  {"layout": "horizontal-bands", "tiers": 3, "flow": "bidirectional"},
  {"layers": [
      {"name": "사용자", "items": ["웹 포털", "모바일 앱", "관리자 콘솔"]},
      {"name": "애플리케이션", "items": ["인증·권한", "업무 로직", "API 게이트웨이", "배치 처리"]},
      {"name": "데이터", "items": ["운영 DB", "캐시", "파일 스토리지"]}],
   "flow": "vertical-bidirectional"},
  ["tier_names", "tier_items", "tier_color", "band_arrows"],
  recommended_use=["시스템 구성", "서비스 구조", "기술 아키텍처"])

E("SVC-003", "service", "플랫폼 구조도 (입력·플랫폼·출력)", F1, 3,
  ["구조도", "플랫폼", "시스템구성", "입출력", "아키텍처"],
  {"layout": "center-hub", "inputs": 3, "outputs": 3, "flow": "left-right"},
  {"layers": [
      {"name": "입력", "items": ["공공데이터", "내부 시스템", "외부 API"]},
      {"name": "플랫폼", "items": ["수집 엔진", "처리 엔진", "저장소", "API 서버"]},
      {"name": "출력", "items": ["Open API", "대시보드", "성과 리포트"]}],
   "flow": "left-right"},
  ["platform_label", "core_modules", "input_modules", "output_modules", "arrows"],
  recommended_use=["서비스 구조", "플랫폼 소개", "시스템 구성"])

E("SVC-004", "service", "데이터 흐름도 (수집→통합→활용)", F1, 4,
  ["구조도", "데이터흐름", "시스템구성", "ETL", "연계"],
  {"layout": "converge-diverge", "sources": 4, "uses": 4, "flow": "left-right"},
  {"layers": [
      {"name": "소스", "items": ["공공데이터 포털", "내부 업무 시스템", "외부 제휴 API", "센서·로그"]},
      {"name": "통합", "items": ["데이터 통합·정제(ETL)"]},
      {"name": "활용", "items": ["정책 의사결정", "대민 서비스", "통계 분석", "연계 기관 제공"]}],
   "flow": "left-right"},
  ["source_nodes", "hub_label", "use_nodes", "connectors"],
  recommended_use=["데이터 흐름", "사업 흐름", "서비스 구조"])

E("SVC-005", "service", "서비스 운영 프로세스 (스윔레인)", F2, 1,
  ["운영프로세스", "스윔레인", "구조도", "흐름도", "업무흐름"],
  {"layout": "swimlane", "lanes": 2, "stages": 4, "flow": "left-right"},
  {"layers": [
      {"name": "신청 기관", "items": ["① 신청서 제출", "④ 서비스 제공"]},
      {"name": "수행 기관", "items": ["② 서류 검토", "③ 승인 심의"]}],
   "flow": "left-right-swimlane"},
  ["lane_names", "stage_labels", "lane_assignment", "connectors"],
  recommended_use=["운영 프로세스", "사업 흐름", "업무 절차"])

E("SVC-006", "service", "사업 추진 개념도 (로직 모델)", F2, 2,
  ["구조도", "로직모델", "개념도", "성과관리", "흐름도"],
  {"layout": "logic-model", "columns": 4, "flow": "left-right"},
  {"layers": [
      {"name": "투입", "items": ["예산·인력", "인프라", "데이터 자산"]},
      {"name": "활동", "items": ["시스템 구축", "데이터 정제", "교육·홍보"]},
      {"name": "산출", "items": ["구축 시스템", "정제 데이터셋", "운영 매뉴얼"]},
      {"name": "성과", "items": ["업무 효율화", "대민 만족도", "성과 관리 체계"]}],
   "flow": "left-right"},
  ["stage_names", "stage_items", "stage_color", "arrows"],
  recommended_use=["사업 흐름", "성과 관리", "추진 체계"])

E("SVC-007", "service", "정보 연계도 (허브-스포크)", F2, 3,
  ["구조도", "연계도", "허브스포크", "기관연계", "데이터흐름"],
  {"layout": "hub-spoke", "spokes": 6, "flow": "bidirectional"},
  {"layers": [
      {"name": "허브", "items": ["데이터 연계 허브"]},
      {"name": "스포크", "items": ["중앙부처", "광역 지자체", "기초 지자체", "공공기관", "유관 기관", "민간 협력사"]}],
   "flow": "hub-radial-bidirectional"},
  ["hub_label", "spoke_labels", "spoke_color", "connectors"],
  recommended_use=["정보 연계", "기관 협력", "서비스 구조"])

E("SVC-008", "service", "통합 관제·모니터링 구조도", F2, 4,
  ["구조도", "모니터링", "관제", "시스템구성", "흐름도"],
  {"layout": "pipeline", "stages": 4, "flow": "left-right"},
  {"layers": [
      {"name": "관제센터", "items": ["통합 관제 센터"]},
      {"name": "파이프라인", "items": ["수집", "분석", "대시보드", "알림"]},
      {"name": "소스", "items": ["서버", "네트워크", "애플리케이션", "사용자 이벤트"]}],
   "flow": "left-right"},
  ["center_label", "stage_labels", "stage_desc", "stage_color", "arrows"],
  recommended_use=["시스템 구성", "운영 모니터링", "서비스 구조"])

frag = c.write_fragment("SVC", entries)

print("FILE1", f1)
print("FILE2", f2)
print("FRAG", frag)
print("ENTRIES", len(entries))
