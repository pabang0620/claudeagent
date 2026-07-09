# -*- coding: utf-8 -*-
"""TML(타임라인/로드맵) 대량 확충 — TML-017~TML-200 (184개).

파라미터화 배치 생성기: 12개 구조 패밀리 × 노드수(3~6) × 색 오프셋(0~3) 스윕.
- 노드수 3~6, 색 오프셋 0~3 = 변형 16종. 12패밀리 × 16 = 192 → 앞 184개 사용(TML-017~200).
- 파일 분산: 패밀리별 1파일 TML_bulk_<slug>_v1.pptx (각 ≤16슬라이드, ≤25 규칙 충족). 1슬라이드 1에셋.
- 다중도형은 c.group_asset 으로 자체완결 grpSp 그룹화(에셋당 1콜 복사). id_caption 은 그룹 밖.
- 색은 c.role / c.C 만. 페이지이미지·SmartArt·pic 금지. 날짜견적 금지(M1/1단계/Q1 상대표기).
- 라벨은 제안서 더미(풀에서 슬라이스, 리터럴 중복 회피).
"""
import sys, math
import os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'lib'))
import common as c
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

R = c.role
CC = c.C

# 코퍼레이트 팔레트 순환(블루/틸/시안/네이비/퍼플)
PALETTE = [R("accent_primary"), R("accent_secondary"), R("accent_point"),
           CC["navy_600"], R("sub_header")]


def col(base_idx, off):
    return PALETTE[(base_idx + off) % len(PALETTE)]


# ---- 제안서 더미 라벨 풀 (리터럴 1회 정의 후 슬라이스) ----
STAGE_NAMES = ["착수·분석", "요구정의", "기본설계", "상세설계", "시스템개발", "통합시험"]
STAGE_FOCUS = ["현황 분석·요구 도출", "과제 정의·범위 확정", "아키텍처·화면 설계",
               "데이터·인터페이스 설계", "핵심 기능 구현", "통합·성능·이관"]
DELIVERABLES = ["착수보고서", "요구정의서", "설계서", "개발산출물", "시험결과서", "완료보고서"]
WHEN_M = ["M1", "M2", "M3", "M4", "M5", "M6"]
WHEN_STEP = ["1단계", "2단계", "3단계", "4단계", "5단계", "6단계"]
PERIODS6 = ["1분기", "2분기", "3분기", "4분기", "5분기", "6분기"]
QUARTERS = ["Q1", "Q2", "Q3", "Q4", "Q5", "Q6"]
TRACK_NAMES = ["기획·관리", "설계·개발", "시험·운영", "확산·교육", "품질·보안", "성과·평가"]
HORIZON = [("Now", "현재 추진"), ("Next", "다음 단계"), ("Later", "향후 과제"),
           ("Future", "확장 과제"), ("Vision", "비전 과제"), ("Beyond", "장기 과제")]
CARDS_POOL = ["핵심 기능 개발", "데이터 이관", "파일럿 운영", "기능 고도화", "연계 확대",
              "성능 최적화", "전면 확산", "자동화·지능화", "지속 운영 체계"]
YEAR_LABELS = ["1년차 기반", "2년차 확산", "3년차 성과", "4년차 정착", "5년차 도약", "6년차 성숙"]
YEAR_FOCUS = ["체계 수립·기반 구축", "기능 확대·고도화", "전면 확산·성과 창출",
              "자립 운영·지속 확산", "고도화·지능화", "지속 성장·확대"]


def title_block(S, slide, text, sub=None):
    S.append(c.add_text(slide, 0.55, 0.52, 12.2, 0.5, text, size=18, bold=True,
                        color=R("header_fill"), align=PP_ALIGN.LEFT, font=c.FONT_H))
    if sub:
        S.append(c.add_text(slide, 0.55, 1.02, 12.2, 0.32, sub, size=11,
                            color=R("muted_text"), align=PP_ALIGN.LEFT))


def BOX(S, slide, *a, **k):
    sp = c.add_box(slide, *a, **k); S.append(sp); return sp


def TXT(S, slide, *a, **k):
    tb = c.add_text(slide, *a, **k); S.append(tb); return tb


def CN(S, slide, *a, **k):
    cn = c.connector(slide, *a, **k); S.append(cn); return cn


# =========================================================
# 1) 수평 타임라인 (가로선 + 마일스톤 원 n + 상하 라벨 교차)
# =========================================================
def horizontal_timeline(slide, aid, n, off):
    S = []
    title_block(S, slide, "수평 타임라인 (%d 마일스톤)" % n,
                "가로 기준선 위 마일스톤을 상·하 라벨로 교차 배치")
    x0, x1, ymid = 1.35, 11.98, 4.0
    CN(S, slide, x0, ymid, x1, ymid, color=R("accent_primary"), w=3.0)
    tip = BOX(S, slide, x1 - 0.02, ymid - 0.16, 0.34, 0.32, fill=R("accent_primary"),
              line=None, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    tip.rotation = 90
    dia = 0.46
    cw = min(2.1, (x1 - x0) / n)
    for i in range(n):
        cx = x0 + (x1 - x0) * (i / (n - 1)) if n > 1 else (x0 + x1) / 2
        color = col(i, off)
        above = (i % 2 == 0)
        ly = ymid - 1.95 if above else ymid + 0.98
        if above:
            CN(S, slide, cx, ymid - dia / 2, cx, ymid - 0.95, color=R("border"), w=1.0)
        else:
            CN(S, slide, cx, ymid + dia / 2, cx, ymid + 0.95, color=R("border"), w=1.0)
        badge = BOX(S, slide, cx - 0.5, ly, 1.0, 0.34, fill=color, line=None,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(badge, WHEN_M[i], size=10, bold=True, color=R("header_text"))
        TXT(S, slide, cx - cw / 2, ly + 0.4, cw, 0.6, STAGE_NAMES[i], size=11, bold=True,
            color=R("body_text"), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
        node = BOX(S, slide, cx - dia / 2, ymid - dia / 2, dia, dia, fill=color,
                   line=CC["white"], line_w=2.5, shape=MSO_SHAPE.OVAL)
        c.set_shape_text(node, str(i + 1), size=13, bold=True, color=R("header_text"))
    c.group_asset(slide, S, aid)


# =========================================================
# 2) 마일스톤 카드 로우 (상단 점 연결선 + 카드 n)
# =========================================================
def milestone_card_row(slide, aid, n, off):
    S = []
    title_block(S, slide, "마일스톤 카드 로우 (%d 카드)" % n,
                "상단 연결선의 마일스톤 점에 카드를 매다는 형태")
    x0, total = 1.1, 11.13
    gap = 0.4 if n <= 4 else 0.28
    cw = (total - gap * (n - 1)) / n
    line_y, card_y, card_h = 2.0, 2.55, 3.75
    CN(S, slide, x0 + cw / 2, line_y, x0 + total - cw / 2, line_y,
       color=R("accent_primary"), w=2.5)
    for i in range(n):
        cx = x0 + i * (cw + gap); center = cx + cw / 2; color = col(i, off)
        BOX(S, slide, center - 0.13, line_y - 0.13, 0.26, 0.26, fill=color,
            line=CC["white"], line_w=2.0, shape=MSO_SHAPE.OVAL)
        CN(S, slide, center, line_y + 0.13, center, card_y, color=R("border"), w=1.0)
        BOX(S, slide, cx, card_y, cw, card_h, fill=CC["gray_050"], line=color, line_w=1.25,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        badge = BOX(S, slide, cx + (cw - 1.05) / 2, card_y + 0.2, 1.05, 0.4, fill=color,
                    line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(badge, WHEN_M[i], size=11, bold=True, color=R("header_text"))
        TXT(S, slide, cx + 0.14, card_y + 0.72, cw - 0.28, 0.5, STAGE_NAMES[i], size=12,
            bold=True, color=R("body_text"), align=PP_ALIGN.CENTER)
        TXT(S, slide, cx + 0.16, card_y + 1.28, cw - 0.32, card_h - 1.45, STAGE_FOCUS[i],
            size=10, color=R("muted_text"), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    c.group_asset(slide, S, aid)


# =========================================================
# 3) 간트 (도형막대 다중 트랙 + 세그먼트)
# =========================================================
def gantt(slide, aid, n, off):
    S = []
    title_block(S, slide, "간트 (%d 트랙 × 기간)" % n,
                "트랙별 도형 막대 세그먼트로 병행 작업을 표현")
    lx, lw = 1.1, 2.5
    gx0, gx1 = lx + lw, 12.2
    gw = gx1 - gx0
    periods = PERIODS6
    npc = len(periods); colw = gw / npc
    hy, hh = 1.7, 0.48
    ry0 = hy + hh
    rh = min(0.7, (6.9 - ry0) / n)
    BOX(S, slide, lx, hy, lw, hh, fill=R("header_fill"), line=CC["white"], line_w=1.0)
    TXT(S, slide, lx, hy, lw, hh, "트랙 / 기간", size=11, bold=True,
        color=R("header_text"), align=PP_ALIGN.CENTER)
    for j, pl in enumerate(periods):
        px = gx0 + j * colw
        BOX(S, slide, px, hy, colw, hh, fill=CC["navy_600"], line=CC["white"], line_w=1.0)
        TXT(S, slide, px, hy, colw, hh, pl, size=10, bold=True,
            color=R("header_text"), align=PP_ALIGN.CENTER)
    for j in range(npc + 1):
        gxj = gx0 + j * colw
        CN(S, slide, gxj, ry0, gxj, ry0 + rh * n, color=CC["gray_300"], w=0.75)
    for i in range(n):
        ry = ry0 + i * rh
        base = R("row_stripe") if i % 2 == 0 else R("row_base")
        BOX(S, slide, lx, ry, lw, rh, fill=base, line=CC["gray_300"], line_w=0.75)
        TXT(S, slide, lx + 0.12, ry, lw - 0.2, rh, TRACK_NAMES[i], size=11, bold=True,
            color=R("body_text"), align=PP_ALIGN.LEFT)
        BOX(S, slide, gx0, ry, gw, rh, fill=base, line=None)
        start = min(i, npc - 2); span = 2 if start + 2 <= npc else 1
        color = col(i, off)
        bar = BOX(S, slide, gx0 + start * colw + 0.06, ry + rh * 0.24,
                  span * colw - 0.12, rh * 0.52, fill=color, line=None,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(bar, STAGE_NAMES[i], size=9, bold=True, color=R("header_text"))
    c.group_asset(slide, S, aid)


# =========================================================
# 4) 단계 로드맵 (셰브론 n + 산출물)
# =========================================================
def stage_chevron(slide, aid, n, off):
    S = []
    title_block(S, slide, "단계 로드맵 (%d 셰브론 + 산출물)" % n,
                "셰브론으로 단계 흐름을 잇고 각 단계 아래 산출물 배치")
    x0, total = 0.7, 11.9
    ov = 0.28
    cw = (total + ov * (n - 1)) / n
    chev_y, chev_h = 1.7, 0.95
    box_y = 3.05; box_h = 3.15
    for i in range(n):
        cx = x0 + i * (cw - ov); color = col(i, off)
        chev = BOX(S, slide, cx, chev_y, cw, chev_h, fill=color, line=CC["white"],
                   line_w=1.5, shape=MSO_SHAPE.CHEVRON)
        c.set_shape_text(chev, "%s  %s" % (WHEN_STEP[i], STAGE_NAMES[i]), size=12,
                         bold=True, color=R("header_text"))
        bx = cx + (ov if i > 0 else 0)
        bw = cw - ov if 0 < i < n - 1 else cw - (ov if i > 0 else 0)
        BOX(S, slide, bx, box_y, bw, box_h, fill=CC["gray_050"], line=color, line_w=1.25,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, bx + 0.08, box_y + 0.1, bw - 0.16, 0.34, "주요 산출물", size=10,
            bold=True, color=color, align=PP_ALIGN.CENTER)
        for k in range(3):
            iy = box_y + 0.56 + k * 0.62
            BOX(S, slide, bx + 0.18, iy + 0.08, 0.1, 0.1, fill=color, line=None,
                shape=MSO_SHAPE.OVAL)
            TXT(S, slide, bx + 0.38, iy - 0.06, bw - 0.5, 0.5, DELIVERABLES[(i + k) % len(DELIVERABLES)],
                size=10, color=R("body_text"), align=PP_ALIGN.LEFT)
    c.group_asset(slide, S, aid)


# =========================================================
# 5) 분기 그리드 (활동군 n행 × 4 분기)
# =========================================================
def quarter_grid(slide, aid, n, off):
    S = []
    title_block(S, slide, "분기 그리드 (%d 활동군 × 분기)" % n,
                "활동군을 행, 분기를 열로 교차한 그리드형 추진일정")
    lx, lw = 1.1, 2.3
    gx0, gx1 = lx + lw, 12.2
    gw = gx1 - gx0
    quarters = QUARTERS[:4]
    nq = len(quarters); colw = gw / nq
    hy, hh = 1.75, 0.55
    ry0 = hy + hh + 0.08
    rh = (6.9 - ry0) / n
    BOX(S, slide, lx, hy, lw, hh, fill=R("header_fill"), line=CC["white"], line_w=1.0)
    TXT(S, slide, lx, hy, lw, hh, "활동군 / 분기", size=11, bold=True,
        color=R("header_text"), align=PP_ALIGN.CENTER)
    for j, q in enumerate(quarters):
        px = gx0 + j * colw
        BOX(S, slide, px, hy, colw, hh, fill=CC["navy_600"], line=CC["white"], line_w=1.0)
        TXT(S, slide, px, hy, colw, hh, q, size=11, bold=True,
            color=R("header_text"), align=PP_ALIGN.CENTER)
    for i in range(n):
        ry = ry0 + i * rh; color = col(i, off)
        base = R("row_stripe") if i % 2 == 0 else R("row_base")
        BOX(S, slide, lx, ry, lw, rh, fill=CC["gray_050"], line=color, line_w=1.5)
        TXT(S, slide, lx + 0.1, ry, lw - 0.2, rh, TRACK_NAMES[i], size=11, bold=True,
            color=color, align=PP_ALIGN.CENTER)
        for j in range(nq):
            px = gx0 + j * colw
            BOX(S, slide, px, ry, colw, rh, fill=base, line=CC["gray_300"], line_w=0.75)
            if (i + j) % 2 == 0:
                pill = BOX(S, slide, px + 0.12, ry + rh * 0.22, colw - 0.24, rh * 0.56,
                           fill=color, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                c.set_shape_text(pill, STAGE_NAMES[(i + j) % len(STAGE_NAMES)], size=9,
                                 bold=False, color=R("header_text"))
    c.group_asset(slide, S, aid)


# =========================================================
# 6) 세로 타임라인 (좌 시간축 + 우 이벤트 카드 n)
# =========================================================
def vertical_timeline(slide, aid, n, off):
    S = []
    title_block(S, slide, "세로 타임라인 (%d 이벤트)" % n,
                "왼쪽 세로 기준선의 시점 노드에 이벤트 카드를 우측 배치")
    axis_x, top, bottom = 3.05, 1.75, 7.0
    rowh = (bottom - top) / n
    CN(S, slide, axis_x, top, axis_x, bottom, color=R("accent_primary"), w=3.0)
    for i in range(n):
        ny = top + rowh * (i + 0.5); color = col(i, off)
        badge = BOX(S, slide, axis_x - 1.75, ny - 0.22, 1.35, 0.44, fill=color, line=None,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(badge, WHEN_STEP[i], size=11, bold=True, color=R("header_text"))
        BOX(S, slide, axis_x - 0.16, ny - 0.16, 0.32, 0.32, fill=color, line=CC["white"],
            line_w=2.0, shape=MSO_SHAPE.OVAL)
        CN(S, slide, axis_x + 0.16, ny, axis_x + 0.55, ny, color=R("border"), w=1.0)
        BOX(S, slide, axis_x + 0.55, ny - rowh * 0.4, 8.65, rowh * 0.8,
            fill=CC["gray_050"], line=color, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, axis_x + 0.75, ny - rowh * 0.4 + 0.04, 8.3, 0.4, STAGE_NAMES[i], size=12,
            bold=True, color=R("body_text"), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
        TXT(S, slide, axis_x + 0.75, ny - rowh * 0.4 + 0.44, 8.3, rowh * 0.8 - 0.48, STAGE_FOCUS[i],
            size=10, color=R("muted_text"), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    c.group_asset(slide, S, aid)


# =========================================================
# 7) 원형 연차 로드맵 (원 둘레 단계 n)
# =========================================================
def circular_roadmap(slide, aid, n, off):
    S = []
    title_block(S, slide, "원형 연차 로드맵 (%d 단계)" % n,
                "중심 허브를 둘러싼 원 둘레에 단계를 등간격 배치")
    cx, cy = 6.55, 4.35
    node_ring, label_ring, node_d = 1.55, 2.75, 0.92
    hub = BOX(S, slide, cx - 0.9, cy - 0.9, 1.8, 1.8, fill=R("header_fill"),
              line=CC["white"], line_w=2.5, shape=MSO_SHAPE.OVAL)
    c.set_shape_text(hub, "추진\n로드맵", size=14, bold=True, color=R("header_text"))
    for i in range(n):
        ang = math.radians(-90 + i * 360.0 / n)
        nx = cx + node_ring * math.cos(ang); ny = cy + node_ring * math.sin(ang)
        lxp = cx + label_ring * math.cos(ang); lyp = cy + label_ring * math.sin(ang)
        color = col(i, off)
        CN(S, slide, cx, cy, nx, ny, color=CC["gray_300"], w=1.25)
        node = BOX(S, slide, nx - node_d / 2, ny - node_d / 2, node_d, node_d, fill=color,
                   line=CC["white"], line_w=2.0, shape=MSO_SHAPE.OVAL)
        c.set_shape_text(node, str(i + 1), size=15, bold=True, color=R("header_text"))
        cw2, ch2 = 2.1, 0.86
        BOX(S, slide, lxp - cw2 / 2, lyp - ch2 / 2, cw2, ch2, fill=CC["gray_050"],
            line=color, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, lxp - cw2 / 2 + 0.06, lyp - ch2 / 2 + 0.06, cw2 - 0.12, 0.34,
            YEAR_LABELS[i], size=11, bold=True, color=color, align=PP_ALIGN.CENTER)
        TXT(S, slide, lxp - cw2 / 2 + 0.06, lyp - ch2 / 2 + 0.42, cw2 - 0.12, 0.36,
            YEAR_FOCUS[i], size=9, color=R("muted_text"), align=PP_ALIGN.CENTER)
    c.group_asset(slide, S, aid)


# =========================================================
# 8) 스윔레인 타임라인 (레인 n × 기간, 기간 걸친 막대)
# =========================================================
def swimlane(slide, aid, n, off):
    S = []
    title_block(S, slide, "스윔레인 타임라인 (%d 레인 × 기간)" % n,
                "레인별 가로 트랙에 기간을 걸친 활동 막대를 배치")
    lx, lw = 1.1, 1.9
    gx0, gx1 = lx + lw, 12.2
    gw = gx1 - gx0
    periods = PERIODS6
    npc = len(periods); colw = gw / npc
    hy, hh = 1.75, 0.5
    ry0 = hy + hh + 0.1
    gap = 0.18
    lane_h = (7.05 - ry0 - gap * (n - 1)) / n
    BOX(S, slide, lx, hy, lw, hh, fill=R("header_fill"), line=CC["white"], line_w=1.0)
    TXT(S, slide, lx, hy, lw, hh, "레인 / 기간", size=11, bold=True,
        color=R("header_text"), align=PP_ALIGN.CENTER)
    for j, pl in enumerate(periods):
        px = gx0 + j * colw
        BOX(S, slide, px, hy, colw, hh, fill=CC["navy_600"], line=CC["white"], line_w=1.0)
        TXT(S, slide, px, hy, colw, hh, pl, size=10, bold=True,
            color=R("header_text"), align=PP_ALIGN.CENTER)
    lanes_bottom = ry0 + lane_h * n + gap * (n - 1)
    for j in range(1, npc):
        gxj = gx0 + j * colw
        CN(S, slide, gxj, ry0, gxj, lanes_bottom, color=CC["gray_300"], w=0.75)
    for i in range(n):
        ly = ry0 + i * (lane_h + gap); color = col(i, off)
        BOX(S, slide, lx, ly, lw, lane_h, fill=CC["gray_050"], line=color, line_w=1.5,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, lx + 0.1, ly, lw - 0.2, lane_h, TRACK_NAMES[i], size=11, bold=True,
            color=color, align=PP_ALIGN.CENTER)
        BOX(S, slide, gx0, ly, gw, lane_h, fill=CC["gray_100"], line=None)
        start = min(i, npc - 2); span = 2 if start + 2 <= npc else 1
        bar = BOX(S, slide, gx0 + start * colw + 0.08, ly + lane_h * 0.24,
                  span * colw - 0.16, lane_h * 0.52, fill=color, line=None,
                  shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(bar, STAGE_NAMES[i], size=9, bold=True, color=R("header_text"))
    c.group_asset(slide, S, aid)


# =========================================================
# 9) 계단 상승 로드맵 (좌하 → 우상 계단 n)
# =========================================================
def staircase(slide, aid, n, off):
    S = []
    title_block(S, slide, "계단 상승 로드맵 (%d 단계)" % n,
                "좌하단에서 우상단으로 상승하는 계단형 성장 로드맵")
    x0, y_bottom = 1.0, 6.6
    total_w = 11.6
    sw = total_w / n
    rise = 0.78
    step_h = 0.9
    for i in range(n):
        color = col(i, off)
        sx = x0 + i * sw
        sy = y_bottom - (i + 1) * rise
        # 계단 라이저(수직 지지 밴드)
        BOX(S, slide, sx, sy, sw - 0.15, y_bottom - sy, fill=CC["gray_100"],
            line=CC["gray_300"], line_w=0.75)
        # 단계 트레드(색 박스)
        tread = BOX(S, slide, sx, sy, sw - 0.15, step_h, fill=color, line=CC["white"],
                    line_w=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(tread, "%s\n%s" % (WHEN_STEP[i], STAGE_NAMES[i]), size=11,
                         bold=True, color=R("header_text"))
        # 초점 라벨
        TXT(S, slide, sx, sy + step_h + 0.05, sw - 0.15, 0.7, STAGE_FOCUS[i], size=9,
            color=R("muted_text"), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    # 상승 화살표(좌하 → 우상)
    CN(S, slide, x0 - 0.1, y_bottom + 0.05, x0 + total_w, y_bottom - n * rise + 0.2,
       color=R("accent_primary"), w=2.0)
    c.group_asset(slide, S, aid)


# =========================================================
# 10) Now/Next/Later (호라이즌 n열)
# =========================================================
def horizon_roadmap(slide, aid, n, off):
    S = []
    title_block(S, slide, "우선순위 호라이즌 로드맵 (%d 구간)" % n,
                "우선순위 기반 구간(Now/Next/Later…) 로드맵")
    x0, total = 1.1, 11.13
    gap = 0.35 if n <= 4 else 0.24
    cw = (total - gap * (n - 1)) / n
    y0, col_h = 1.75, 4.95
    ncards = 3 if n <= 4 else 2
    for i in range(n):
        cx = x0 + i * (cw + gap); color = col(i, off)
        label, tag = HORIZON[i]
        hd = BOX(S, slide, cx, y0, cw, 0.95, fill=color, line=None,
                 shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        tf = hd.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label; c.set_kfont(r, c.FONT_H, 16, True, R("header_text"))
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = tag; c.set_kfont(r2, c.FONT_B, 9, False, R("header_text"))
        BOX(S, slide, cx, y0 + 0.95, cw, col_h - 0.95, fill=CC["gray_050"],
            line=R("border"), line_w=1.0)
        for k in range(ncards):
            ky = y0 + 1.15 + k * 0.92
            cb = BOX(S, slide, cx + 0.18, ky, cw - 0.36, 0.74, fill=R("row_base"),
                     line=color, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            c.set_shape_text(cb, CARDS_POOL[(i + k) % len(CARDS_POOL)], size=10,
                             bold=False, color=R("body_text"))
    c.group_asset(slide, S, aid)


# =========================================================
# 11) 플래그 마일스톤 (기준선 위 깃발 n)
# =========================================================
def flag_milestone(slide, aid, n, off):
    S = []
    title_block(S, slide, "플래그 마일스톤 (%d 깃발)" % n,
                "가로 기준선의 시점마다 깃대와 깃발로 마일스톤을 강조")
    x0, x1 = 1.3, 12.0
    base_y = 5.9
    CN(S, slide, x0, base_y, x1, base_y, color=R("accent_primary"), w=3.0)
    pole_h = 2.9
    fw, fh = 1.55, 0.66
    for i in range(n):
        cx = x0 + (x1 - x0) * (i / (n - 1)) if n > 1 else (x0 + x1) / 2
        color = col(i, off)
        # 기준선 노드
        BOX(S, slide, cx - 0.13, base_y - 0.13, 0.26, 0.26, fill=color, line=CC["white"],
            line_w=2.0, shape=MSO_SHAPE.OVAL)
        # 깃대
        CN(S, slide, cx, base_y, cx, base_y - pole_h, color=R("muted_text"), w=1.5)
        # 깃발(깃대 상단 오른쪽)
        flag = BOX(S, slide, cx + 0.02, base_y - pole_h, fw, fh, fill=color, line=None,
                   shape=MSO_SHAPE.PENTAGON)
        c.set_shape_text(flag, WHEN_M[i], size=11, bold=True, color=R("header_text"))
        # 라벨(깃발 아래)
        TXT(S, slide, cx - 0.15, base_y - pole_h + fh + 0.04, fw + 0.3, 0.55, STAGE_NAMES[i],
            size=10, bold=True, color=R("body_text"), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    c.group_asset(slide, S, aid)


# =========================================================
# 12) 3트랙 로드맵 (트랙 3 × 페이즈 n열)
# =========================================================
def track3_roadmap(slide, aid, n, off):
    S = []
    title_block(S, slide, "3트랙 로드맵 (3 트랙 × %d 페이즈)" % n,
                "3개 병렬 트랙을 페이즈 열로 정렬한 로드맵")
    lx, lw = 1.1, 2.2
    gx0, gx1 = lx + lw, 12.2
    gw = gx1 - gx0
    colw = gw / n
    hy, hh = 1.75, 0.55
    ry0 = hy + hh + 0.1
    ntr = 3
    rh = (6.85 - ry0) / ntr
    # 페이즈 헤더
    BOX(S, slide, lx, hy, lw, hh, fill=R("header_fill"), line=CC["white"], line_w=1.0)
    TXT(S, slide, lx, hy, lw, hh, "트랙 / 페이즈", size=11, bold=True,
        color=R("header_text"), align=PP_ALIGN.CENTER)
    for j in range(n):
        px = gx0 + j * colw
        BOX(S, slide, px, hy, colw, hh, fill=CC["navy_600"], line=CC["white"], line_w=1.0)
        TXT(S, slide, px, hy, colw, hh, WHEN_STEP[j], size=10, bold=True,
            color=R("header_text"), align=PP_ALIGN.CENTER)
    for t in range(ntr):
        ry = ry0 + t * rh; color = col(t, off)
        BOX(S, slide, lx, ry, lw, rh, fill=CC["gray_050"], line=color, line_w=1.5,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, lx + 0.1, ry, lw - 0.2, rh, TRACK_NAMES[t], size=12, bold=True,
            color=color, align=PP_ALIGN.CENTER)
        # 트랙 진행 막대(페이즈 걸침)
        BOX(S, slide, gx0, ry + rh * 0.5 - 0.02, gw, 0.04, fill=CC["gray_300"], line=None)
        for j in range(n):
            px = gx0 + j * colw
            node = BOX(S, slide, px + colw / 2 - 0.6, ry + rh * 0.2, 1.2, rh * 0.6,
                       fill=color, line=CC["white"], line_w=1.5,
                       shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            c.set_shape_text(node, STAGE_NAMES[(t + j) % len(STAGE_NAMES)], size=9,
                             bold=True, color=R("header_text"))
    c.group_asset(slide, S, aid)


# =========================================================
# 패밀리 메타 (slug, fn, 이름, 단위, 방향, 태그, 추천용도, 바인딩키)
# =========================================================
FAMILIES = [
    dict(slug="htimeline", fn=horizontal_timeline, name="수평 타임라인", unit="마일스톤",
         orient="horizontal", tags=["타임라인", "수평", "마일스톤", "추진일정"],
         use=["추진일정", "주요 마일스톤", "연혁"], binds=["when-label", "node-text", "count", "color"]),
    dict(slug="mcardrow", fn=milestone_card_row, name="마일스톤 카드 로우", unit="카드",
         orient="horizontal", tags=["타임라인", "마일스톤", "카드", "점연결", "추진일정"],
         use=["추진일정", "주요 마일스톤", "로드맵"], binds=["when-label", "card-head", "card-desc", "count", "color"]),
    dict(slug="gantt", fn=gantt, name="간트 (도형막대)", unit="트랙",
         orient="horizontal", tags=["타임라인", "간트", "다중트랙", "막대", "추진일정"],
         use=["추진일정", "병행 작업 일정", "로드맵"], binds=["track-name", "bar-text", "bar-length", "count", "color"]),
    dict(slug="stagechevron", fn=stage_chevron, name="단계 셰브론 로드맵", unit="단계",
         orient="horizontal", tags=["로드맵", "단계별", "셰브론", "산출물", "추진일정"],
         use=["단계별 계획", "산출물 매핑", "로드맵"], binds=["stage-name", "deliverable-text", "count", "color"]),
    dict(slug="qgrid", fn=quarter_grid, name="분기 그리드", unit="활동군",
         orient="grid", tags=["로드맵", "분기별", "그리드", "활동군", "추진일정"],
         use=["분기 계획", "활동 매핑", "추진일정"], binds=["cat-name", "cell-text", "count", "color"]),
    dict(slug="vtimeline", fn=vertical_timeline, name="세로 타임라인", unit="이벤트",
         orient="vertical", tags=["타임라인", "세로", "시간축", "이벤트카드", "추진일정"],
         use=["추진일정", "단계별 계획", "연혁"], binds=["when-label", "event-head", "event-desc", "count", "color"]),
    dict(slug="circular", fn=circular_roadmap, name="원형 연차 로드맵", unit="단계",
         orient="radial", tags=["로드맵", "원형", "연차별", "허브", "추진일정"],
         use=["중장기 계획", "연차 로드맵", "발전 방향"], binds=["stage-label", "focus-text", "count", "color"]),
    dict(slug="swimlane", fn=swimlane, name="스윔레인 타임라인", unit="레인",
         orient="horizontal", tags=["타임라인", "스윔레인", "레인", "기간막대", "추진일정"],
         use=["추진일정", "병행 작업 일정", "로드맵"], binds=["lane-name", "bar-text", "bar-length", "count", "color"]),
    dict(slug="staircase", fn=staircase, name="계단 상승 로드맵", unit="단계",
         orient="ascending", tags=["로드맵", "계단", "상승", "성장", "단계별"],
         use=["성장 로드맵", "단계별 발전", "중장기 계획"], binds=["stage-name", "focus-text", "count", "color"]),
    dict(slug="horizon", fn=horizon_roadmap, name="Now/Next/Later 로드맵", unit="구간",
         orient="horizontal", tags=["로드맵", "우선순위", "Now-Next-Later", "호라이즌", "제품로드맵"],
         use=["우선순위 로드맵", "제품 로드맵", "추진 방향"], binds=["col-label", "card-text", "count", "color"]),
    dict(slug="flag", fn=flag_milestone, name="플래그 마일스톤", unit="깃발",
         orient="horizontal", tags=["타임라인", "마일스톤", "깃발", "플래그", "추진일정"],
         use=["추진일정", "주요 마일스톤", "핵심 이벤트"], binds=["when-label", "flag-text", "count", "color"]),
    dict(slug="track3", fn=track3_roadmap, name="3트랙 로드맵", unit="페이즈",
         orient="horizontal", tags=["로드맵", "3트랙", "병렬트랙", "페이즈", "추진일정"],
         use=["병렬 트랙 로드맵", "다영역 계획", "추진일정"], binds=["track-name", "phase-text", "count", "color"]),
]

# =========================================================
# 빌드 — 12패밀리 × 변형16(노드3~6 × 색0~3) 스윕, 앞 184개
# =========================================================
FILE_TPL = "decks/05_timeline/TML_bulk_%s_v1.pptx"
START_NUM = 17
TARGET = 184

# 변형 인덱스 v: node = 3 + v%4, offset = v//4 (노드 완전 스윕 × 색 오프셋 0~3)
combos = []
for v in range(16):
    node = 3 + (v % 4)
    off = v // 4
    for fam_idx in range(len(FAMILIES)):
        combos.append((fam_idx, node, off))
combos = combos[:TARGET]

decks = {}          # slug -> Presentation
deck_slides = {}    # slug -> 현재 슬라이드 수
entries = []

for k, (fam_idx, node, off) in enumerate(combos):
    fam = FAMILIES[fam_idx]
    slug = fam["slug"]
    aid = "TML-%03d" % (START_NUM + k)
    if slug not in decks:
        decks[slug] = c.new_deck(); deck_slides[slug] = 0
    prs = decks[slug]
    s = c.blank_slide(prs)
    c.id_caption(s, aid)
    fam["fn"](s, aid, node, off)
    deck_slides[slug] += 1
    slide_idx = deck_slides[slug]
    file_rel = FILE_TPL % slug
    entries.append(c.entry(
        aid, "TML",
        "%s (%d%s)" % (fam["name"], node, fam["unit"]),
        file_rel, slide_idx,
        fam["tags"] + ["노드%d" % node],
        {"nodes": node, "color_offset": off, "family": slug},
        {"count": node, "orientation": fam["orient"], "palette_offset": off, "unit": fam["unit"]},
        fam["binds"],
        recommended_use=fam["use"]))

# 파일 저장
saved = []
for slug, prs in decks.items():
    path = c.save_deck(prs, FILE_TPL % slug)
    saved.append((slug, path, deck_slides[slug]))

frag = c.write_fragment("TML_bulk", entries)

print("ENTRIES", len(entries))
print("FIRST", entries[0]["id"], "LAST", entries[-1]["id"])
print("FILES", len(saved))
for slug, path, cnt in saved:
    print("  %-14s %2d slides  %s" % (slug, cnt, path))
print("FRAG", frag)
