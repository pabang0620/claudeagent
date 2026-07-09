# -*- coding: utf-8 -*-
"""TBL_gov — 정부보고서(gov) 트랙 표 인포그래픽 6종. TBL-201 ~ TBL-206.

캔버스: gov 비율 11.93×8.50in (EMU 10905360×7772400) — prs.slide_width/height 직접 override.
색·폰트: design-tokens.json 최상위 gov_theme(color/role/font/size_pt) 만 참조. 매직 헥스 금지.
병합(gridSpan/rowSpan)은 python-pptx cell.merge()로 생성 시점에 굽는다(BUILD_SPEC 병합표 원칙).
표는 graphicFrame 단독 앵커 — 그룹 불필요, gf.name = "asset:<ID>" 만 지정.
출력: decks/01_tables/TBL_gov_v1.pptx (신규 파일). 매니페스트: _incoming/manifest_TBL_gov.json (고유 프래그먼트).
"""
import sys
sys.path.insert(0, '/home/pabang/myapp/.claude/pptx-asset-library/generators/lib')
import common as c
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------- gov_theme 로드 (additive 최상위 키, common.TOKENS 경유) ----------
GOV = c.TOKENS["gov_theme"]
GC = {k: RGBColor.from_string(v) for k, v in GOV["color"].items()}
def GR(name):  # gov role -> RGBColor
    return GC[GOV["role"][name]]
GOV_FONT = GOV["font"]
GOV_SZ = GOV["size_pt"]
F_HEAD = GOV_FONT["heading"]["typeface"]
F_SUB = GOV_FONT["subheading"]["typeface"]
F_BODY = GOV_FONT["body"]["typeface"]
F_CAP = GOV_FONT["caption"]["typeface"]
F_LABEL = GOV_FONT["label"]["typeface"]

GOV_W = Emu(10905360)
GOV_H = Emu(7772400)
GOV_W_IN = 10905360 / 914400.0
GOV_H_IN = 7772400 / 914400.0

NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid

L = PP_ALIGN.LEFT; Rg = PP_ALIGN.RIGHT; Ct = PP_ALIGN.CENTER


def new_gov_deck():
    prs = c.new_deck()
    prs.slide_width = GOV_W
    prs.slide_height = GOV_H
    return prs


def title_block(slide, text, sub=None):
    c.add_text(slide, 0.5, 0.35, GOV_W_IN - 1.0, 0.5, text, size=GOV_SZ["heading"],
               bold=True, color=GR("accent_primary"), align=L, font=F_HEAD)
    if sub:
        c.add_text(slide, 0.5, 0.85, GOV_W_IN - 1.0, 0.32, sub, size=GOV_SZ["caption"],
                   color=GC["shape_fill_gray"], align=L, font=F_CAP)


def add_table(slide, x, y, w, h, nrows, ncols, col_widths=None, row_h=None):
    gf = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    sid = tblPr.find(qn('a:tableStyleId'))
    if sid is None:
        sid = etree.SubElement(tblPr, qn('a:tableStyleId'))
    sid.text = NO_STYLE
    if col_widths:
        for j, cw in enumerate(col_widths):
            tbl.columns[j].width = Inches(cw)
    if row_h:
        for i, rh in enumerate(row_h):
            tbl.rows[i].height = Inches(rh)
    return gf, tbl


def set_cell(cell, text, size=11, bold=False, color=None, fill=None,
             align=Ct, font=None, ml=5):
    if fill is not None:
        cell.fill.solid(); cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid(); cell.fill.fore_color.rgb = GC["bg_primary"]
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Pt(ml)
    cell.margin_top = cell.margin_bottom = Pt(2)
    tf = cell.text_frame; tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = line
        c.set_kfont(r, font or F_BODY, size, bold, color or GC["text"])


def cell_borders(cell, spec):
    tcPr = cell._tc.get_or_add_tcPr()
    for edge, tag in (('B', 'a:lnB'), ('T', 'a:lnT'), ('R', 'a:lnR'), ('L', 'a:lnL')):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        if edge not in spec:
            continue
        color, wpt = spec[edge]
        ln = etree.Element(qn(tag))
        ln.set('w', str(int(wpt * 12700))); ln.set('cap', 'flat')
        if color is None:
            etree.SubElement(ln, qn('a:noFill'))
        else:
            sf = etree.SubElement(ln, qn('a:solidFill'))
            etree.SubElement(sf, qn('a:srgbClr')).set('val', str(color))
        tcPr.insert(0, ln)


def grid(tbl, color, wpt=None):
    wpt = wpt or GOV["geom"]["border_pt"]
    for row in tbl.rows:
        for cell in row.cells:
            cell_borders(cell, {'L': (color, wpt), 'R': (color, wpt),
                                'T': (color, wpt), 'B': (color, wpt)})


# 그리드선 전용 중립색 — gov_theme에 별도 border/gray 토큰이 없어 표헤더색(8F99AF)을
# 저굵기(border_pt)로 재사용한다(룰5: 새 헥스 금지, 기존 토큰 재사용).
GRID_LINE = GR("header_fill")


# =========================================================
# TBL-201  색코딩 병합헤더 비교표 (colspan+rowspan, 연도 비교)
# =========================================================
def t201(slide):
    title_block(slide, "연도별 사업 실행성과 비교표", "색코딩 병합헤더 · 2025년 실적 vs 2026년 계획")
    x, y = 0.5, 1.35
    TW = GOV_W_IN - 1.0
    cw = [TW * 0.24, TW * 0.13, TW * 0.13, TW * 0.13, TW * 0.13, TW * 0.24]
    nrows = 6  # 2 header + 4 body
    rh = 0.55
    gf, tbl = add_table(slide, x, y, TW, rh * nrows, nrows, 6, col_widths=cw, row_h=[rh] * nrows)

    # 헤더 병합
    tbl.cell(0, 0).merge(tbl.cell(1, 0))
    tbl.cell(0, 1).merge(tbl.cell(0, 2))
    tbl.cell(0, 3).merge(tbl.cell(0, 4))
    tbl.cell(0, 5).merge(tbl.cell(1, 5))

    set_cell(tbl.cell(0, 0), "구분", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), font=F_SUB)
    set_cell(tbl.cell(0, 1), "2025년 실적", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill_2"), font=F_SUB)
    set_cell(tbl.cell(0, 3), "2026년 계획", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill_2"), font=F_SUB)
    set_cell(tbl.cell(0, 5), "증감", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), font=F_SUB)
    for j, hd in zip((1, 2, 3, 4), ("목표", "실적", "목표", "실적")):
        set_cell(tbl.cell(1, j), hd, size=11, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill"), font=F_LABEL)

    rows = [
        ("사업 추진율", "80%", "76%", "100%", "94%", "▲18%p", True),
        ("예산 집행률", "70%", "68%", "95%", "91%", "▲23%p", True),
        ("콘텐츠 제작 건수", "120건", "108건", "150건", "121건", "▲13건", True),
        ("고객 만족도(CSAT)", "85점", "82점", "90점", "79점", "▼3점", False),
    ]
    for i, (lab, t1, a1, t2, a2, delta, up) in enumerate(rows, start=2):
        base = GC["bg_panel"] if i % 2 == 0 else GC["bg_primary"]
        set_cell(tbl.cell(i, 0), lab, size=11, bold=True, fill=base, align=L)
        for j, v in zip((1, 2, 3, 4), (t1, a1, t2, a2)):
            set_cell(tbl.cell(i, j), v, size=11, fill=base, align=Rg)
        dcol = C_TEAL if up else GR("warn")
        set_cell(tbl.cell(i, 5), delta, size=11, bold=True, color=dcol, fill=base, align=Ct)
    grid(tbl, GRID_LINE)
    gf.name = "asset:TBL-201"
    c.id_caption(slide, "TBL-201")


# =========================================================
# TBL-202  카테고리 구분(보라) 2단헤더 표
# =========================================================
def t202(slide):
    title_block(slide, "추진체계 카테고리 구분표", "2단 헤더 · 카테고리(보라) + 세부 지표")
    x, y = 0.5, 1.35
    TW = GOV_W_IN - 1.0
    cw = [TW * 0.13] + [TW * 0.87 / 6] * 6
    nrows = 6  # 2 header + 4 body
    rh = 0.55
    gf, tbl = add_table(slide, x, y, TW, rh * nrows, nrows, 7, col_widths=cw, row_h=[rh] * nrows)

    tbl.cell(0, 0).merge(tbl.cell(1, 0))
    tbl.cell(0, 1).merge(tbl.cell(0, 2))
    tbl.cell(0, 3).merge(tbl.cell(0, 4))
    tbl.cell(0, 5).merge(tbl.cell(0, 6))

    set_cell(tbl.cell(0, 0), "구분", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), font=F_SUB)
    for j, hd in zip((1, 3, 5), ("추진체계", "성과관리", "환류체계")):
        set_cell(tbl.cell(0, j), hd, size=12, bold=True, color=GC["bg_primary"],
                 fill=GR("category"), font=F_SUB)
    subs = ["조직", "역할", "지표", "목표치", "점검주기", "개선조치"]
    for j, hd in zip(range(1, 7), subs):
        set_cell(tbl.cell(1, j), hd, size=11, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill"), font=F_LABEL)

    rows = [
        ("총괄", "사업총괄반", "예산·일정 총괄", "사업추진율", "100%", "월 1회", "지연시 즉시 보고"),
        ("실행", "실행분과", "세부과제 수행", "과제완료율", "95%", "월 1회", "지연과제 재계획"),
        ("자문", "전문가 자문단", "기술·정책 자문", "자문충족도", "90%", "분기 1회", "자문의견 반영률 관리"),
        ("평가", "성과평가위원회", "성과 검증·환류", "환류이행률", "100%", "반기 1회", "미이행 과제 재조치"),
    ]
    for i, row in enumerate(rows, start=2):
        base = GC["bg_panel"] if i % 2 == 0 else GC["bg_primary"]
        set_cell(tbl.cell(i, 0), row[0], size=11, bold=True, fill=base, align=Ct)
        for j, v in enumerate(row[1:], start=1):
            set_cell(tbl.cell(i, j), v, size=10, fill=base, align=L if j in (1, 5) else Ct)
    grid(tbl, GRID_LINE)
    gf.name = "asset:TBL-202"
    c.id_caption(slide, "TBL-202")


# =========================================================
# TBL-203  상태 신호등(traffic-light) 표
# =========================================================
def t203(slide):
    title_block(slide, "과제 진행상태 점검표", "상태 신호등 셀 · 완료/진행중/지연/대기 4단계")
    x, y = 0.5, 1.35
    TW = GOV_W_IN - 1.0
    cw = [TW * 0.08, TW * 0.32, TW * 0.18, TW * 0.16, TW * 0.26]
    nrows = 7  # 1 header + 6 body
    rh = 0.55
    gf, tbl = add_table(slide, x, y, TW, rh * nrows, nrows, 5, col_widths=cw, row_h=[rh] * nrows)

    heads = ["No", "과제명", "담당", "진행상태", "비고"]
    for j, hd in enumerate(heads):
        set_cell(tbl.cell(0, j), hd, size=12, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill"), font=F_SUB, align=L if j == 1 else Ct)

    # 신호등 색: gov_theme엔 성공(초록) 토큰이 없어 상위 top-level teal_500을 성공색으로
    # 차용(룰5: 대체 토큰 사용, 신규 헥스 금지). 나머지는 gov_theme 자체 토큰.
    signal = {"완료": C_TEAL, "진행중": GR("accent_primary"),
              "지연": GR("warn"), "대기": GC["shape_fill_gray"]}
    rows = [
        (1, "시스템 구축", "개발운영팀", "완료", "검수 완료"),
        (2, "데이터 연계", "정보화팀", "완료", "연계테스트 통과"),
        (3, "콘텐츠 제작", "사업운영팀", "진행중", "68% 진척"),
        (4, "대국민 홍보", "홍보팀", "진행중", "채널 3종 운영"),
        (5, "성과평가 체계 구축", "평가관리팀", "지연", "일정 재조정 필요"),
        (6, "유지보수 체계 수립", "운영지원팀", "대기", "착수 전"),
    ]
    for i, (no, name, owner, status, note) in enumerate(rows, start=1):
        base = GC["bg_panel"] if i % 2 == 0 else GC["bg_primary"]
        set_cell(tbl.cell(i, 0), no, size=11, fill=base, align=Ct)
        set_cell(tbl.cell(i, 1), name, size=11, bold=True, fill=base, align=L)
        set_cell(tbl.cell(i, 2), owner, size=10, fill=base, align=Ct)
        set_cell(tbl.cell(i, 3), status, size=11, bold=True, color=GC["bg_primary"],
                 fill=signal[status], align=Ct)
        set_cell(tbl.cell(i, 4), note, size=10, fill=base, align=L)
    grid(tbl, GRID_LINE)
    gf.name = "asset:TBL-203"
    c.id_caption(slide, "TBL-203")


# =========================================================
# TBL-204  평가/배점표 (rowspan 카테고리 + 합계행)
# =========================================================
def t204(slide):
    title_block(slide, "정성·정량 평가 배점표", "평가영역 행병합(rowSpan) · 하단 합계행 열병합(gridSpan)")
    x, y = 0.5, 1.35
    TW = GOV_W_IN - 1.0
    cw = [TW * 0.22, TW * 0.46, TW * 0.16, TW * 0.16]
    nrows = 8  # 1 header + 6 body + 1 total
    rh = 0.55
    gf, tbl = add_table(slide, x, y, TW, rh * nrows, nrows, 4, col_widths=cw, row_h=[rh] * nrows)

    heads = ["평가영역", "세부 평가지표", "배점", "획득점수"]
    for j, hd in enumerate(heads):
        set_cell(tbl.cell(0, j), hd, size=12, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill"), font=F_SUB, align=L if j == 1 else Ct)

    # (row_idx_start, span, 라벨, [(세부지표, 배점, 획득점수), ...])
    groups = [
        (1, 2, "사업 이해도", [("사업목표 부합성", "10", "9"), ("추진전략 타당성", "10", "8")]),
        (3, 3, "수행 역량", [("조직 및 인력", "10", "9"), ("유사사업 실적", "10", "7"), ("협력체계", "10", "8")]),
        (6, 1, "가격", [("예산 적정성", "10", "9")]),
    ]
    for start, span, label, items in groups:
        if span > 1:
            tbl.cell(start, 0).merge(tbl.cell(start + span - 1, 0))
        base_lab = GR("header_fill_2")
        set_cell(tbl.cell(start, 0), label, size=11, bold=True, color=GC["bg_primary"],
                 fill=base_lab, align=Ct)
        for k, (crit, pt, sc) in enumerate(items):
            i = start + k
            base = GC["bg_panel"] if i % 2 == 0 else GC["bg_primary"]
            set_cell(tbl.cell(i, 1), crit, size=10, fill=base, align=L)
            set_cell(tbl.cell(i, 2), pt, size=11, fill=base, align=Ct)
            set_cell(tbl.cell(i, 3), sc, size=11, bold=True, color=GR("accent_primary"),
                     fill=base, align=Ct)
    # 합계행
    li = nrows - 1
    tbl.cell(li, 0).merge(tbl.cell(li, 1))
    set_cell(tbl.cell(li, 0), "합계", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), align=Ct)
    set_cell(tbl.cell(li, 2), "60", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), align=Ct)
    set_cell(tbl.cell(li, 3), "50", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), align=Ct)
    grid(tbl, GRID_LINE)
    gf.name = "asset:TBL-204"
    c.id_caption(slide, "TBL-204")


# =========================================================
# TBL-205  예산/집행표 (rowspan 카테고리 + 합계행)
# =========================================================
def t205(slide):
    title_block(slide, "사업비 예산·집행 현황표", "구분 행병합(rowSpan) · 하단 합계행 열병합(gridSpan)")
    x, y = 0.5, 1.35
    TW = GOV_W_IN - 1.0
    cw = [TW * 0.14, TW * 0.24, TW * 0.16, TW * 0.16, TW * 0.14, TW * 0.16]
    nrows = 7  # 1 header + 5 body + 1 total
    rh = 0.55
    gf, tbl = add_table(slide, x, y, TW, rh * nrows, nrows, 6, col_widths=cw, row_h=[rh] * nrows)

    heads = ["구분", "세부항목", "예산액", "집행액", "집행률", "비고"]
    for j, hd in enumerate(heads):
        set_cell(tbl.cell(0, j), hd, size=12, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill"), font=F_SUB, align=L if j in (1, 5) else Ct)

    groups = [
        (1, 2, "인건비", [("정규인력", "90,000", "82,400", "91.6%", "정상집행"),
                          ("전문인력", "40,000", "35,100", "87.8%", "정상집행")]),
        (3, 3, "사업비", [("콘텐츠 제작", "120,000", "98,500", "82.1%", "정상집행"),
                          ("시스템 구축", "150,000", "112,300", "74.9%", "순연 집행"),
                          ("홍보·확산", "60,000", "51,200", "85.3%", "정상집행")]),
    ]
    for start, span, label, items in groups:
        if span > 1:
            tbl.cell(start, 0).merge(tbl.cell(start + span - 1, 0))
        set_cell(tbl.cell(start, 0), label, size=11, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill_2"), align=Ct)
        for k, (item, bud, exe, rate, note) in enumerate(items):
            i = start + k
            base = GC["bg_panel"] if i % 2 == 0 else GC["bg_primary"]
            set_cell(tbl.cell(i, 1), item, size=10, fill=base, align=L)
            set_cell(tbl.cell(i, 2), bud, size=10, fill=base, align=Rg)
            set_cell(tbl.cell(i, 3), exe, size=10, bold=True, color=GR("accent_primary"),
                     fill=base, align=Rg)
            set_cell(tbl.cell(i, 4), rate, size=10, bold=True, fill=base, align=Ct)
            set_cell(tbl.cell(i, 5), note, size=9, fill=base, align=L)
    li = nrows - 1
    tbl.cell(li, 0).merge(tbl.cell(li, 1))
    set_cell(tbl.cell(li, 0), "합계", size=12, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), align=Ct)
    set_cell(tbl.cell(li, 2), "460,000", size=11, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), align=Rg)
    set_cell(tbl.cell(li, 3), "379,500", size=11, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), align=Rg)
    set_cell(tbl.cell(li, 4), "82.5%", size=11, bold=True, color=GC["bg_primary"],
             fill=GR("header_fill"), align=Ct)
    set_cell(tbl.cell(li, 5), "-", size=10, color=GC["bg_primary"], fill=GR("header_fill"), align=Ct)
    grid(tbl, GRID_LINE)
    gf.name = "asset:TBL-205"
    c.id_caption(slide, "TBL-205")


# =========================================================
# TBL-206  체크리스트표 (rowspan 카테고리 그룹 + 체크칩)
# =========================================================
def t206(slide):
    title_block(slide, "분야별 점검 체크리스트표", "구분 행병합(rowSpan) · 확인여부 칩 셀")
    x, y = 0.5, 1.35
    TW = GOV_W_IN - 1.0
    cw = [TW * 0.14, TW * 0.42, TW * 0.2, TW * 0.24]
    nrows = 7  # 1 header + 6 body
    rh = 0.55
    gf, tbl = add_table(slide, x, y, TW, rh * nrows, nrows, 4, col_widths=cw, row_h=[rh] * nrows)

    heads = ["구분", "점검항목", "확인여부", "비고"]
    for j, hd in enumerate(heads):
        set_cell(tbl.cell(0, j), hd, size=12, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill"), font=F_SUB, align=L if j in (1, 3) else Ct)

    groups = [
        (1, 2, "보안", [("취약점 점검 조치", "완료", "최종 조치완료"),
                        ("개인정보 처리 점검", "완료", "DPIA 완료")]),
        (3, 2, "성능", [("부하테스트", "진행중", "3차 테스트 중"),
                        ("장애대응 체계", "예정", "매뉴얼 작성 예정")]),
        (5, 2, "운영", [("운영 매뉴얼 수립", "완료", "배포 완료"),
                        ("교육·이관", "예정", "착수 전")]),
    ]
    status_col = {"완료": C_TEAL, "진행중": GR("accent_primary"), "예정": GC["shape_fill_gray"]}
    for start, span, label, items in groups:
        tbl.cell(start, 0).merge(tbl.cell(start + span - 1, 0))
        set_cell(tbl.cell(start, 0), label, size=11, bold=True, color=GC["bg_primary"],
                 fill=GR("header_fill_2"), align=Ct)
        for k, (item, status, note) in enumerate(items):
            i = start + k
            base = GC["bg_panel"] if i % 2 == 0 else GC["bg_primary"]
            set_cell(tbl.cell(i, 1), item, size=11, fill=base, align=L)
            chip = "✓ %s" % status if status == "완료" else status
            set_cell(tbl.cell(i, 2), chip, size=11, bold=True, color=GC["bg_primary"],
                     fill=status_col[status], align=Ct)
            set_cell(tbl.cell(i, 3), note, size=10, fill=base, align=L)
    grid(tbl, GRID_LINE)
    gf.name = "asset:TBL-206"
    c.id_caption(slide, "TBL-206")


# ---- teal 성공색(top-level 대체 토큰, gov_theme에 성공색 부재) ----
C_TEAL = c.C["teal_500"]


# =========================================================
# 빌드
# =========================================================
def build():
    prs = new_gov_deck()
    fns = (t201, t202, t203, t204, t205, t206)
    for fn in fns:
        fn(c.blank_slide(prs))
    out_rel = "decks/01_tables/TBL_gov_v1.pptx"
    path = c.save_deck(prs, out_rel)

    E = c.entry
    entries = [
        E("TBL-201", "TBL", "gov 연도별 실행성과 비교표 (색코딩 병합헤더)", out_rel, 1,
          ["table", "gov", "merged-header", "comparison", "정부", "병합헤더", "비교"],
          {"cols": 6, "rows": 6, "merged": True}, ["구분", "연도값", "증감"],
          ["header", "row_label", "cell_value"],
          recommended_use=["실행계획서", "연도별 성과비교", "정부보고서"], master="gov"),
        E("TBL-202", "TBL", "gov 추진체계 카테고리 구분표 (보라 2단헤더)", out_rel, 2,
          ["table", "gov", "category", "purple", "정부", "카테고리", "2단헤더"],
          {"cols": 7, "rows": 6, "merged": True}, ["구분", "카테고리", "세부지표"],
          ["header", "category_header", "sub_header", "cell_value"],
          recommended_use=["추진체계", "역할분담", "정부보고서"], master="gov"),
        E("TBL-203", "TBL", "gov 과제 진행상태 신호등표", out_rel, 3,
          ["table", "gov", "status", "traffic-light", "정부", "신호등", "진행상태"],
          {"cols": 5, "rows": 7}, ["과제명", "담당", "상태"],
          ["header", "item", "status_cell"],
          recommended_use=["진행상태 점검", "과제관리", "정부보고서"], master="gov"),
        E("TBL-204", "TBL", "gov 정성·정량 평가 배점표 (행병합+합계행)", out_rel, 4,
          ["table", "gov", "evaluation", "score", "정부", "평가", "배점"],
          {"cols": 4, "rows": 8, "merged": True}, ["평가영역", "세부지표", "배점", "획득점수"],
          ["header", "eval_group", "criteria", "score", "total_row"],
          recommended_use=["평가배점표", "심사기준", "정부보고서"], master="gov"),
        E("TBL-205", "TBL", "gov 예산·집행 현황표 (행병합+합계행)", out_rel, 5,
          ["table", "gov", "budget", "execution", "정부", "예산", "집행"],
          {"cols": 6, "rows": 7, "merged": True}, ["구분", "세부항목", "예산액", "집행액", "집행률"],
          ["header", "budget_group", "item", "amount", "total_row"],
          recommended_use=["예산집행현황", "사업비 보고", "정부보고서"], master="gov"),
        E("TBL-206", "TBL", "gov 분야별 점검 체크리스트표 (행병합+체크칩)", out_rel, 6,
          ["table", "gov", "checklist", "check", "정부", "체크리스트", "점검"],
          {"cols": 4, "rows": 7, "merged": True}, ["구분", "점검항목", "확인여부"],
          ["header", "check_group", "item", "status_chip"],
          recommended_use=["점검표", "완료여부 확인", "정부보고서"], master="gov"),
    ]
    frag = c.write_fragment("TBL_gov", entries)
    print("SAVED:", path)
    print("FRAGMENT:", frag)


if __name__ == "__main__":
    build()
