# -*- coding: utf-8 -*-
"""TML 카테고리 변형팩 v2 (타임라인/로드맵) 에셋 생성 — TML-009~TML-016.

v1과 다른 형태로 8종 신규:
  TML_timeline_v2.pptx : TML-009 스윔레인 / TML-010 마일스톤카드row / TML-013 세로타임라인 / TML-014 간트v2
  TML_roadmap_v2.pptx  : TML-011 원형연차 / TML-012 단계+산출물 / TML-015 분기그리드 / TML-016 Now/Next/Later

규칙: 네이티브 도형(사각/원/오각/화살표)+커넥터만. 페이지이미지·SmartArt 금지.
다중도형은 c.group_asset(slide, [도형들], 'TML-0NN') 으로 그룹핑(에셋당 1콜 복사).
색은 c.role / c.C 만 사용. 날짜견적 금지(M1/1단계/Q1 상대표기). 라벨은 제안서 더미.
"""
import sys, math
sys.path.insert(0, '/home/pabang/myapp/.claude/pptx-asset-library/generators/lib')
import common as c
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

R = c.role
CC = c.C

# 코퍼레이트 팔레트 순환(네이비/블루/틸/시안/그레이)
PALETTE = [R("accent_primary"), R("accent_secondary"), R("accent_point"),
           CC["navy_600"], R("sub_header")]


# ---------------------------------------------------------
# 공통 헬퍼 — 생성 도형을 리스트 S 에 모아 group_asset 으로 묶는다.
# id_caption 은 S 에 넣지 않는다(라이브러리 라벨).
# ---------------------------------------------------------
def title_block(S, slide, text, sub=None):
    S.append(c.add_text(slide, 0.55, 0.52, 12.2, 0.5, text, size=18, bold=True,
                        color=R("header_fill"), align=PP_ALIGN.LEFT, font=c.FONT_H))
    if sub:
        S.append(c.add_text(slide, 0.55, 1.02, 12.2, 0.32, sub, size=11,
                            color=R("muted_text"), align=PP_ALIGN.LEFT))


def BOX(S, slide, *a, **k):
    sp = c.add_box(slide, *a, **k)
    S.append(sp)
    return sp


def TXT(S, slide, *a, **k):
    tb = c.add_text(slide, *a, **k)
    S.append(tb)
    return tb


def CN(S, slide, *a, **k):
    cn = c.connector(slide, *a, **k)
    S.append(cn)
    return cn


# =========================================================
# TML-009 스윔레인 타임라인 (2~3 레인 × 기간, 기간 걸친 막대)
# =========================================================
def swimlane(slide, asset_id, title, sub, periods, lanes):
    """lanes: list of dict {name, color_idx, bars:[(label, start, span)]}."""
    S = []
    title_block(S, slide, title, sub)
    lx, lw = 1.1, 1.9
    gx0 = lx + lw
    gx1 = 12.2
    gw = gx1 - gx0
    npc = len(periods)
    colw = gw / npc
    hy, hh = 1.75, 0.5
    ry0 = hy + hh + 0.1
    lane_h = 1.35
    gap = 0.2
    # 헤더(레인/기간 라벨)
    BOX(S, slide, lx, hy, lw, hh, fill=R("header_fill"), line=CC["white"], line_w=1.0)
    TXT(S, slide, lx, hy, lw, hh, "레인 / 기간", size=11, bold=True,
        color=R("header_text"), align=PP_ALIGN.CENTER)
    for j, pl in enumerate(periods):
        px = gx0 + j * colw
        BOX(S, slide, px, hy, colw, hh, fill=CC["navy_600"], line=CC["white"], line_w=1.0)
        TXT(S, slide, px, hy, colw, hh, pl, size=10, bold=True,
            color=R("header_text"), align=PP_ALIGN.CENTER)
    lanes_bottom = ry0 + lane_h * len(lanes) + gap * (len(lanes) - 1)
    # 기간 구분 세로 그리드
    for j in range(1, npc):
        gxj = gx0 + j * colw
        CN(S, slide, gxj, ry0, gxj, lanes_bottom, color=CC["gray_300"], w=0.75)
    # 레인
    for ti, lane in enumerate(lanes):
        ly = ry0 + ti * (lane_h + gap)
        color = PALETTE[lane["color_idx"] % len(PALETTE)]
        # 레인명 박스
        BOX(S, slide, lx, ly, lw, lane_h, fill=CC["gray_050"], line=color, line_w=1.5,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, lx + 0.1, ly, lw - 0.2, lane_h, lane["name"], size=12, bold=True,
            color=color, align=PP_ALIGN.CENTER)
        # 레인 배경 밴드
        BOX(S, slide, gx0, ly, gw, lane_h, fill=CC["gray_100"], line=None)
        # 기간 걸친 막대
        for (blabel, start, span) in lane["bars"]:
            bx = gx0 + start * colw + 0.08
            bw = span * colw - 0.16
            bar = BOX(S, slide, bx, ly + lane_h * 0.24, bw, lane_h * 0.52, fill=color,
                      line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            c.set_shape_text(bar, blabel, size=9, bold=True, color=R("header_text"))
    c.group_asset(slide, S, asset_id)


# =========================================================
# TML-010 마일스톤 카드 row (상단 점 연결선 + 카드 4개)
# =========================================================
def milestone_card_row(slide, asset_id, title, sub, cards):
    """cards: list of dict {when, head, desc, color_idx}."""
    S = []
    title_block(S, slide, title, sub)
    x0 = 1.1
    total = 11.13
    n = len(cards)
    gap = 0.4
    cw = (total - gap * (n - 1)) / n
    line_y = 2.0
    card_y = 2.55
    card_h = 3.75
    # 상단 연결선
    CN(S, slide, x0 + cw / 2, line_y, x0 + total - cw / 2, line_y,
       color=R("accent_primary"), w=2.5)
    for i, cd in enumerate(cards):
        cx = x0 + i * (cw + gap)
        center = cx + cw / 2
        color = PALETTE[cd["color_idx"] % len(PALETTE)]
        # 연결선 위 점
        BOX(S, slide, center - 0.13, line_y - 0.13, 0.26, 0.26, fill=color,
            line=CC["white"], line_w=2.0, shape=MSO_SHAPE.OVAL)
        # 점 → 카드 리더선
        CN(S, slide, center, line_y + 0.13, center, card_y, color=R("border"), w=1.0)
        # 카드
        BOX(S, slide, cx, card_y, cw, card_h, fill=CC["gray_050"], line=color, line_w=1.25,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        # when 뱃지
        badge = BOX(S, slide, cx + (cw - 1.05) / 2, card_y + 0.2, 1.05, 0.4, fill=color,
                    line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(badge, cd["when"], size=11, bold=True, color=R("header_text"))
        # 제목
        TXT(S, slide, cx + 0.14, card_y + 0.72, cw - 0.28, 0.5, cd["head"], size=12,
            bold=True, color=R("body_text"), align=PP_ALIGN.CENTER)
        # 설명
        TXT(S, slide, cx + 0.16, card_y + 1.28, cw - 0.32, card_h - 1.45, cd["desc"],
            size=10, color=R("muted_text"), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)
    c.group_asset(slide, S, asset_id)


# =========================================================
# TML-011 원형 연차 로드맵 (원 둘레 단계 배치)
# =========================================================
def circular_roadmap(slide, asset_id, title, sub, stages):
    """stages: list of dict {label, focus, color_idx}. 원 둘레 등간격 배치."""
    S = []
    title_block(S, slide, title, sub)
    cx, cy = 6.55, 4.35
    node_ring = 1.55
    label_ring = 2.75
    node_d = 0.92
    # 중심 원
    hub = BOX(S, slide, cx - 0.9, cy - 0.9, 1.8, 1.8, fill=R("header_fill"),
              line=CC["white"], line_w=2.5, shape=MSO_SHAPE.OVAL)
    c.set_shape_text(hub, "추진\n로드맵", size=14, bold=True, color=R("header_text"))
    n = len(stages)
    for i, st in enumerate(stages):
        ang = math.radians(-90 + i * 360.0 / n)
        nx = cx + node_ring * math.cos(ang)
        ny = cy + node_ring * math.sin(ang)
        lxp = cx + label_ring * math.cos(ang)
        lyp = cy + label_ring * math.sin(ang)
        color = PALETTE[st["color_idx"] % len(PALETTE)]
        # 스포크(중심 → 노드)
        CN(S, slide, cx, cy, nx, ny, color=CC["gray_300"], w=1.25)
        # 단계 노드 원
        node = BOX(S, slide, nx - node_d / 2, ny - node_d / 2, node_d, node_d, fill=color,
                   line=CC["white"], line_w=2.0, shape=MSO_SHAPE.OVAL)
        c.set_shape_text(node, str(i + 1), size=15, bold=True, color=R("header_text"))
        # 라벨 카드(바깥쪽)
        cw2, ch2 = 2.1, 0.86
        BOX(S, slide, lxp - cw2 / 2, lyp - ch2 / 2, cw2, ch2, fill=CC["gray_050"],
            line=color, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, lxp - cw2 / 2 + 0.06, lyp - ch2 / 2 + 0.06, cw2 - 0.12, 0.34,
            st["label"], size=11, bold=True, color=color, align=PP_ALIGN.CENTER)
        TXT(S, slide, lxp - cw2 / 2 + 0.06, lyp - ch2 / 2 + 0.42, cw2 - 0.12, 0.36,
            st["focus"], size=9, color=R("muted_text"), align=PP_ALIGN.CENTER)
    c.group_asset(slide, S, asset_id)


# =========================================================
# TML-012 단계 + 산출물 로드맵 (단계 원 → 산출물 박스)
# =========================================================
def stage_deliverable(slide, asset_id, title, sub, stages):
    """stages: list of dict {no, name, deliverables:[...], color_idx}."""
    S = []
    title_block(S, slide, title, sub)
    x0 = 1.1
    total = 11.13
    n = len(stages)
    gap = 0.3
    cw = (total - gap * (n - 1)) / n
    circle_y = 1.75
    circle_d = 1.05
    box_y = 3.35
    box_h = 3.05
    for i, st in enumerate(stages):
        cx = x0 + i * (cw + gap)
        center = cx + cw / 2
        color = PALETTE[st["color_idx"] % len(PALETTE)]
        # 단계 원(번호)
        circ = BOX(S, slide, center - circle_d / 2, circle_y, circle_d, circle_d, fill=color,
                   line=CC["white"], line_w=2.5, shape=MSO_SHAPE.OVAL)
        c.set_shape_text(circ, "%s단계" % st["no"], size=13, bold=True, color=R("header_text"))
        # 단계명
        TXT(S, slide, cx, circle_y + circle_d + 0.06, cw, 0.4, st["name"], size=12,
            bold=True, color=color, align=PP_ALIGN.CENTER)
        # 단계 → 다음 단계 진행 화살표
        if i < n - 1:
            CN(S, slide, center + circle_d / 2, circle_y + circle_d / 2,
               center + cw + gap - circle_d / 2, circle_y + circle_d / 2,
               color=R("accent_primary"), w=1.75)
        # 원 → 산출물 박스 리더선
        CN(S, slide, center, circle_y + circle_d + 0.46, center, box_y, color=R("border"), w=1.0)
        # 산출물 박스
        BOX(S, slide, cx, box_y, cw, box_h, fill=CC["gray_050"], line=color, line_w=1.25,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, cx + 0.08, box_y + 0.1, cw - 0.16, 0.34, "주요 산출물", size=10,
            bold=True, color=color, align=PP_ALIGN.CENTER)
        for k, d in enumerate(st["deliverables"]):
            iy = box_y + 0.56 + k * 0.56
            BOX(S, slide, cx + 0.2, iy + 0.1, 0.1, 0.1, fill=color, line=None, shape=MSO_SHAPE.OVAL)
            TXT(S, slide, cx + 0.4, iy - 0.04, cw - 0.55, 0.5, d, size=10,
                color=R("body_text"), align=PP_ALIGN.LEFT)
    c.group_asset(slide, S, asset_id)


# =========================================================
# TML-013 세로 타임라인 (좌 시간축 + 우 이벤트 카드)
# =========================================================
def vertical_timeline(slide, asset_id, title, sub, events):
    """events: list of dict {when, head, desc, color_idx}."""
    S = []
    title_block(S, slide, title, sub)
    axis_x = 3.05
    top = 1.75
    bottom = 7.0
    n = len(events)
    rowh = (bottom - top) / n
    # 세로 기준선
    CN(S, slide, axis_x, top, axis_x, bottom, color=R("accent_primary"), w=3.0)
    for i, ev in enumerate(events):
        ny = top + rowh * (i + 0.5)
        color = PALETTE[ev["color_idx"] % len(PALETTE)]
        # when 뱃지(축 왼쪽)
        badge = BOX(S, slide, axis_x - 1.75, ny - 0.22, 1.35, 0.44, fill=color, line=None,
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        c.set_shape_text(badge, ev["when"], size=11, bold=True, color=R("header_text"))
        # 축 위 노드
        BOX(S, slide, axis_x - 0.16, ny - 0.16, 0.32, 0.32, fill=color, line=CC["white"],
            line_w=2.0, shape=MSO_SHAPE.OVAL)
        # 노드 → 카드 리더선
        CN(S, slide, axis_x + 0.16, ny, axis_x + 0.55, ny, color=R("border"), w=1.0)
        # 이벤트 카드(축 오른쪽)
        card = BOX(S, slide, axis_x + 0.55, ny - rowh * 0.4, 8.65, rowh * 0.8,
                   fill=CC["gray_050"], line=color, line_w=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        TXT(S, slide, axis_x + 0.75, ny - rowh * 0.4 + 0.06, 8.3, 0.4, ev["head"], size=12,
            bold=True, color=R("body_text"), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
        TXT(S, slide, axis_x + 0.75, ny - rowh * 0.4 + 0.46, 8.3, rowh * 0.8 - 0.5, ev["desc"],
            size=10, color=R("muted_text"), align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP)
    c.group_asset(slide, S, asset_id)


# =========================================================
# TML-014 간트 v2 (색 구분 다중 트랙 + 세그먼트 + 범례)
# =========================================================
def gantt_v2(slide, asset_id, title, sub, periods, tracks, legend):
    """tracks: list of dict {name, segments:[(label, start, span, color_idx)]}.
    legend: list of (name, color_idx)."""
    S = []
    title_block(S, slide, title, sub)
    lx, lw = 1.1, 2.5
    gx0 = lx + lw
    gx1 = 12.2
    gw = gx1 - gx0
    npc = len(periods)
    colw = gw / npc
    hy, hh = 1.7, 0.48
    ry0 = hy + hh
    rh = 0.7
    nrows = len(tracks)
    # 헤더
    BOX(S, slide, lx, hy, lw, hh, fill=R("header_fill"), line=CC["white"], line_w=1.0)
    TXT(S, slide, lx, hy, lw, hh, "트랙 / 기간", size=11, bold=True,
        color=R("header_text"), align=PP_ALIGN.CENTER)
    for j, pl in enumerate(periods):
        px = gx0 + j * colw
        BOX(S, slide, px, hy, colw, hh, fill=CC["navy_600"], line=CC["white"], line_w=1.0)
        TXT(S, slide, px, hy, colw, hh, pl, size=10, bold=True,
            color=R("header_text"), align=PP_ALIGN.CENTER)
    # 세로 그리드
    for j in range(npc + 1):
        gxj = gx0 + j * colw
        CN(S, slide, gxj, ry0, gxj, ry0 + rh * nrows, color=CC["gray_300"], w=0.75)
    # 트랙 행 + 다색 세그먼트
    for i, tr in enumerate(tracks):
        ry = ry0 + i * rh
        base = R("row_stripe") if i % 2 == 0 else R("row_base")
        BOX(S, slide, lx, ry, lw, rh, fill=base, line=CC["gray_300"], line_w=0.75)
        TXT(S, slide, lx + 0.12, ry, lw - 0.2, rh, tr["name"], size=11, bold=True,
            color=R("body_text"), align=PP_ALIGN.LEFT)
        BOX(S, slide, gx0, ry, gw, rh, fill=base, line=None)
        for (slabel, start, span, ci) in tr["segments"]:
            bx = gx0 + start * colw + 0.06
            bw = span * colw - 0.12
            bar = BOX(S, slide, bx, ry + rh * 0.24, bw, rh * 0.52,
                      fill=PALETTE[ci % len(PALETTE)], line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            c.set_shape_text(bar, slabel, size=9, bold=True, color=R("header_text"))
    # 범례
    leg_y = ry0 + rh * nrows + 0.22
    legx = lx
    for (lname, ci) in legend:
        BOX(S, slide, legx, leg_y, 0.28, 0.28, fill=PALETTE[ci % len(PALETTE)],
            line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        tb = TXT(S, slide, legx + 0.36, leg_y - 0.06, 2.4, 0.4, lname, size=10,
                 color=R("body_text"), align=PP_ALIGN.LEFT)
        legx += 0.36 + 0.12 + min(2.4, 0.3 + len(lname) * 0.16) + 0.35
    c.group_asset(slide, S, asset_id)


# =========================================================
# TML-015 분기 그리드 (좌 활동군 열 + 4 분기 열 × 활동 셀)
# =========================================================
def quarter_grid(slide, asset_id, title, sub, quarters, rows):
    """quarters: 헤더 라벨(4). rows: list of dict {cat, cells:[4], color_idx}."""
    S = []
    title_block(S, slide, title, sub)
    lx, lw = 1.1, 2.3
    gx0 = lx + lw
    gx1 = 12.2
    gw = gx1 - gx0
    nq = len(quarters)
    colw = gw / nq
    hy, hh = 1.75, 0.55
    ry0 = hy + hh + 0.08
    nrows = len(rows)
    rh = (6.9 - ry0) / nrows
    # 헤더
    BOX(S, slide, lx, hy, lw, hh, fill=R("header_fill"), line=CC["white"], line_w=1.0)
    TXT(S, slide, lx, hy, lw, hh, "활동군 / 분기", size=11, bold=True,
        color=R("header_text"), align=PP_ALIGN.CENTER)
    for j, q in enumerate(quarters):
        px = gx0 + j * colw
        BOX(S, slide, px, hy, colw, hh, fill=CC["navy_600"], line=CC["white"], line_w=1.0)
        TXT(S, slide, px, hy, colw, hh, q, size=11, bold=True,
            color=R("header_text"), align=PP_ALIGN.CENTER)
    # 행
    for i, row in enumerate(rows):
        ry = ry0 + i * rh
        color = PALETTE[row["color_idx"] % len(PALETTE)]
        base = R("row_stripe") if i % 2 == 0 else R("row_base")
        # 활동군 라벨
        BOX(S, slide, lx, ry, lw, rh, fill=CC["gray_050"], line=color, line_w=1.5)
        TXT(S, slide, lx + 0.1, ry, lw - 0.2, rh, row["cat"], size=11, bold=True,
            color=color, align=PP_ALIGN.CENTER)
        # 분기 셀
        for j in range(nq):
            px = gx0 + j * colw
            BOX(S, slide, px, ry, colw, rh, fill=base, line=CC["gray_300"], line_w=0.75)
            txt = row["cells"][j] if j < len(row["cells"]) else ""
            if txt:
                pill = BOX(S, slide, px + 0.12, ry + rh * 0.2, colw - 0.24, rh * 0.6, fill=color,
                           line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                c.set_shape_text(pill, txt, size=9, bold=False, color=R("header_text"))
    c.group_asset(slide, S, asset_id)


# =========================================================
# TML-016 Now / Next / Later 3열 로드맵
# =========================================================
def now_next_later(slide, asset_id, title, sub, cols):
    """cols: list of dict {label, tag, cards:[...], color_idx} (보통 3열)."""
    S = []
    title_block(S, slide, title, sub)
    x0 = 1.1
    total = 11.13
    n = len(cols)
    gap = 0.35
    cw = (total - gap * (n - 1)) / n
    y0 = 1.75
    col_h = 4.85
    for i, col in enumerate(cols):
        cx = x0 + i * (cw + gap)
        color = PALETTE[col["color_idx"] % len(PALETTE)]
        # 헤더
        hd = BOX(S, slide, cx, y0, cw, 0.95, fill=color, line=None,
                 shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        tf = hd.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = col["label"]; c.set_kfont(r, c.FONT_H, 16, True, R("header_text"))
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = col["tag"]; c.set_kfont(r2, c.FONT_B, 9, False, R("header_text"))
        # 바디
        BOX(S, slide, cx, y0 + 0.95, cw, col_h - 0.95, fill=CC["gray_050"],
            line=R("border"), line_w=1.0)
        # 카드 스택
        for k, card in enumerate(col["cards"]):
            ky = y0 + 1.15 + k * 0.92
            cb = BOX(S, slide, cx + 0.18, ky, cw - 0.36, 0.74, fill=R("row_base"),
                     line=color, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            c.set_shape_text(cb, card, size=10, bold=False, color=R("body_text"))
    c.group_asset(slide, S, asset_id)


# =========================================================
# 빌드
# =========================================================
entries = []


def E(*a, **k):
    entries.append(c.entry(*a, **k))


F1 = "decks/05_timeline/TML_timeline_v2.pptx"
F2 = "decks/05_timeline/TML_roadmap_v2.pptx"

# ---- FILE 1: 타임라인 계열 (TML-009, 010, 013, 014) ----
p1 = c.new_deck()

s = c.blank_slide(p1); c.id_caption(s, "TML-009")
swimlane(s, "TML-009", "스윔레인 타임라인 (3 레인 × 기간)",
         "레인별 가로 트랙에 기간을 걸친 활동 막대를 배치",
         ["1분기", "2분기", "3분기", "4분기", "5분기", "6분기"],
         [{"name": "기획", "color_idx": 0,
           "bars": [("요구 분석", 0, 2), ("과제 정의", 2, 1)]},
          {"name": "구축", "color_idx": 1,
           "bars": [("아키텍처 설계", 1, 2), ("시스템 개발", 3, 2), ("통합 시험", 5, 1)]},
          {"name": "운영", "color_idx": 2,
           "bars": [("시범 운영", 4, 1), ("정식 운영·이관", 5, 1)]}])

s = c.blank_slide(p1); c.id_caption(s, "TML-010")
milestone_card_row(s, "TML-010", "마일스톤 카드 로우 (상단 점 연결)",
                   "상단 연결선과 마일스톤 점에 카드 4개를 매다는 형태",
                   [{"when": "M1", "head": "사업 착수", "color_idx": 0,
                     "desc": "착수보고 · 추진체계 확정 · 킥오프"},
                    {"when": "M2", "head": "설계 완료", "color_idx": 1,
                     "desc": "기본·상세 설계 확정 · 화면 정의"},
                    {"when": "M4", "head": "중간보고", "color_idx": 2,
                     "desc": "핵심 기능 개발 · 단위 시험 완료"},
                    {"when": "M6", "head": "최종납품", "color_idx": 3,
                     "desc": "통합 시험 · 안정화 · 산출물 이관"}])

s = c.blank_slide(p1); c.id_caption(s, "TML-013")
vertical_timeline(s, "TML-013", "세로 타임라인 (좌 시간축 + 우 이벤트)",
                  "왼쪽 세로 기준선의 시점 노드에 이벤트 카드를 우측 배치",
                  [{"when": "1단계", "head": "착수·분석", "color_idx": 0,
                    "desc": "현황 분석 · 요구사항 정의 · 착수보고"},
                   {"when": "2단계", "head": "설계", "color_idx": 1,
                    "desc": "기본·상세 설계 · 화면/데이터 정의서 작성"},
                   {"when": "3단계", "head": "구축", "color_idx": 2,
                    "desc": "시스템 개발 · 단위 시험 · 중간보고"},
                   {"when": "4단계", "head": "시험·안정화", "color_idx": 3,
                    "desc": "통합 시험 · 안정화 · 이관 교육 · 최종납품"}])

s = c.blank_slide(p1); c.id_caption(s, "TML-014")
gantt_v2(s, "TML-014", "간트 v2 (색 구분 다중 트랙)",
         "트랙별 다색 세그먼트로 병행 작업을 구분하고 범례 제공",
         ["1분기", "2분기", "3분기", "4분기", "5분기", "6분기"],
         [{"name": "기획·분석", "segments": [("요구", 0, 2, 0), ("정의", 2, 1, 0)]},
          {"name": "설계", "segments": [("기본", 1, 1, 1), ("상세", 2, 2, 1)]},
          {"name": "개발", "segments": [("코어", 2, 2, 2), ("확장", 4, 1, 2)]},
          {"name": "시험·이관", "segments": [("통합", 4, 1, 3), ("안정화", 5, 1, 3)]}],
         [("기획·분석", 0), ("설계", 1), ("개발", 2), ("시험·이관", 3)])

f1 = c.save_deck(p1, F1)

# ---- FILE 2: 로드맵 계열 (TML-011, 012, 015, 016) ----
p2 = c.new_deck()

s = c.blank_slide(p2); c.id_caption(s, "TML-011")
circular_roadmap(s, "TML-011", "원형 연차 로드맵 (원 둘레 단계)",
                 "중심 허브를 둘러싼 원 둘레에 단계를 등간격 배치",
                 [{"label": "1년차 기반", "focus": "체계 수립·기반 구축", "color_idx": 0},
                  {"label": "2년차 확산", "focus": "기능 확대·고도화", "color_idx": 1},
                  {"label": "3년차 성과", "focus": "전면 확산·성과 창출", "color_idx": 2},
                  {"label": "4년차 정착", "focus": "자립 운영·지속 확산", "color_idx": 3}])

s = c.blank_slide(p2); c.id_caption(s, "TML-012")
stage_deliverable(s, "TML-012", "단계 + 산출물 로드맵 (단계 원 → 산출물)",
                  "단계 원을 화살표로 잇고 각 단계 아래 산출물 박스 배치",
                  [{"no": 1, "name": "착수·분석", "color_idx": 0,
                    "deliverables": ["착수보고서", "현황 분석서", "요구사항 정의서"]},
                   {"no": 2, "name": "설계", "color_idx": 1,
                    "deliverables": ["기본 설계서", "상세 설계서", "화면 정의서"]},
                   {"no": 3, "name": "구축", "color_idx": 2,
                    "deliverables": ["개발 산출물", "단위 시험서", "중간보고서"]},
                   {"no": 4, "name": "시험·안정화", "color_idx": 3,
                    "deliverables": ["통합 시험서", "안정화 보고", "완료보고서"]}])

s = c.blank_slide(p2); c.id_caption(s, "TML-015")
quarter_grid(s, "TML-015", "분기 그리드 (활동군 × 4 분기)",
             "활동군을 행으로, 분기를 열로 교차한 그리드형 추진일정",
             ["Q1", "Q2", "Q3", "Q4"],
             [{"cat": "기획·관리", "color_idx": 0,
               "cells": ["착수·현황분석", "요구정의", "중간점검", "완료보고"]},
              {"cat": "설계·개발", "color_idx": 1,
               "cells": ["", "기본·상세설계", "시스템 개발", "고도화 개발"]},
              {"cat": "시험·운영", "color_idx": 2,
               "cells": ["", "", "단위 시험", "통합·안정화"]},
              {"cat": "확산·교육", "color_idx": 3,
               "cells": ["", "", "매뉴얼 작성", "이관 교육"]}])

s = c.blank_slide(p2); c.id_caption(s, "TML-016")
now_next_later(s, "TML-016", "Now / Next / Later 3열 로드맵",
               "우선순위 기반 3구간(현재·다음·향후) 로드맵",
               [{"label": "Now", "tag": "현재 추진", "color_idx": 0,
                 "cards": ["핵심 기능 개발", "데이터 이관", "파일럿 운영"]},
                {"label": "Next", "tag": "다음 단계", "color_idx": 1,
                 "cards": ["기능 고도화", "연계 시스템 확대", "성능 최적화"]},
                {"label": "Later", "tag": "향후 과제", "color_idx": 2,
                 "cards": ["전면 확산", "자동화·지능화", "지속 운영 체계"]}])

f2 = c.save_deck(p2, F2)

# =========================================================
# 매니페스트 엔트리
# =========================================================
E("TML-009", "timeline", "스윔레인 타임라인 (3 레인 × 기간)", F1, 1,
  ["타임라인", "스윔레인", "레인", "기간막대", "추진일정"],
  {"lanes": 3, "periods": 6, "swimlane": True},
  {"periods": ["1분기", "2분기", "3분기", "4분기", "5분기", "6분기"],
   "lanes": [{"name": "기획"}, {"name": "구축"}, {"name": "운영"}],
   "count": 3, "orientation": "horizontal"},
  ["lane-name", "bar-text", "bar-length", "count", "color"],
  recommended_use=["추진일정", "병행 작업 일정", "로드맵"])

E("TML-010", "timeline", "마일스톤 카드 로우 (상단 점 연결)", F1, 2,
  ["타임라인", "마일스톤", "카드", "점연결", "추진일정"],
  {"cards": 4, "top_connector": True},
  {"milestones": [{"when": "M1", "head": "사업 착수"}, {"when": "M2", "head": "설계 완료"},
                  {"when": "M4", "head": "중간보고"}, {"when": "M6", "head": "최종납품"}],
   "count": 4, "orientation": "horizontal"},
  ["when-label", "card-head", "card-desc", "count", "color"],
  recommended_use=["추진일정", "주요 마일스톤", "로드맵"])

E("TML-011", "timeline", "원형 연차 로드맵 (원 둘레 단계)", F2, 1,
  ["로드맵", "원형", "연차별", "허브", "추진일정"],
  {"stages": 4, "circular": True},
  {"stages": [{"label": "1년차 기반"}, {"label": "2년차 확산"},
              {"label": "3년차 성과"}, {"label": "4년차 정착"}],
   "count": 4, "orientation": "radial"},
  ["stage-label", "focus-text", "count", "color"],
  recommended_use=["중장기 계획", "연차 로드맵", "발전 방향"])

E("TML-012", "timeline", "단계 + 산출물 로드맵 (단계 원 → 산출물)", F2, 2,
  ["로드맵", "단계별", "산출물", "화살표", "추진일정"],
  {"stages": 4, "deliverables_per_stage": 3},
  {"stages": [{"no": 1, "name": "착수·분석"}, {"no": 2, "name": "설계"},
              {"no": 3, "name": "구축"}, {"no": 4, "name": "시험·안정화"}],
   "count": 4, "orientation": "horizontal"},
  ["stage-name", "deliverable-text", "count", "color"],
  recommended_use=["단계별 계획", "산출물 매핑", "로드맵"])

E("TML-013", "timeline", "세로 타임라인 (좌 시간축 + 우 이벤트)", F1, 3,
  ["타임라인", "세로", "시간축", "이벤트카드", "추진일정"],
  {"events": 4, "vertical": True},
  {"events": [{"when": "1단계", "head": "착수·분석"}, {"when": "2단계", "head": "설계"},
              {"when": "3단계", "head": "구축"}, {"when": "4단계", "head": "시험·안정화"}],
   "count": 4, "orientation": "vertical"},
  ["when-label", "event-head", "event-desc", "count", "color"],
  recommended_use=["추진일정", "단계별 계획", "연혁"])

E("TML-014", "timeline", "간트 v2 (색 구분 다중 트랙)", F1, 4,
  ["타임라인", "간트", "다중트랙", "세그먼트", "범례", "추진일정"],
  {"tracks": 4, "periods": 6, "multi_segment": True, "legend": True},
  {"periods": ["1분기", "2분기", "3분기", "4분기", "5분기", "6분기"],
   "tracks": [{"name": "기획·분석"}, {"name": "설계"}, {"name": "개발"}, {"name": "시험·이관"}],
   "count": 4, "orientation": "horizontal"},
  ["track-name", "segment-text", "bar-length", "count", "color"],
  recommended_use=["추진일정", "병행 작업 일정", "로드맵"])

E("TML-015", "timeline", "분기 그리드 (활동군 × 4 분기)", F2, 3,
  ["로드맵", "분기별", "그리드", "활동군", "추진일정"],
  {"rows": 4, "quarters": 4, "grid": True},
  {"quarters": ["Q1", "Q2", "Q3", "Q4"],
   "rows": [{"cat": "기획·관리"}, {"cat": "설계·개발"}, {"cat": "시험·운영"}, {"cat": "확산·교육"}],
   "count": 4, "orientation": "grid"},
  ["cat-name", "cell-text", "count", "color"],
  recommended_use=["분기 계획", "활동 매핑", "추진일정"])

E("TML-016", "timeline", "Now / Next / Later 3열 로드맵", F2, 4,
  ["로드맵", "우선순위", "Now-Next-Later", "3열", "제품로드맵"],
  {"cols": 3, "priority_roadmap": True},
  {"cols": [{"label": "Now", "tag": "현재 추진"}, {"label": "Next", "tag": "다음 단계"},
            {"label": "Later", "tag": "향후 과제"}],
   "count": 3, "orientation": "horizontal"},
  ["col-label", "card-text", "count", "color"],
  recommended_use=["우선순위 로드맵", "제품 로드맵", "추진 방향"])

frag = c.write_fragment("TML_v2", entries)

print("FILE1", f1)
print("FILE2", f2)
print("FRAG", frag)
print("ENTRIES", len(entries))
